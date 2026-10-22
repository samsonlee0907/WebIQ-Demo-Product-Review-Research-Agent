"""Build an original, editable site-research discussion deck without network calls.

The reference presentation informed geometry, typography and colors only.
No reference text, logos, images or protected Office parts are reproduced.
Authoring dependency: python-pptx.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
from urllib.parse import urlsplit

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
NAVY, DEEP, CYAN = "0B1E3F", "12305E", "50E6FF"
BLUE, PURPLE, TEAL = "0078D4", "5C2D91", "007A6E"
INK, GRAY, MUTED = "201F1E", "605E5C", "8A8886"
WHITE, BORDER, PALE = "FFFFFF", "DDE4EE", "F3F7FC"
FONT, STRONG, MONO = "Segoe UI", "Segoe UI Semibold", "Consolas"
METHODS = (("webiq_browse", "WebIQ Browse"), ("live", "Azure BAT Tool"),
           ("client_browser", "Local Playwright"), ("vm_playwright", "Azure VM + Playwright"))
PORTAL = "http://127.0.0.1:8011"
SOURCES = {
    "S1": ("WebIQ Browse API", "https://learn.microsoft.com/en-us/web-iq/reference/browse"),
    "S2": ("WebIQ Web Search API", "https://learn.microsoft.com/en-us/web-iq/reference/web-search"),
    "S3": ("WebIQ FAQ", "https://learn.microsoft.com/en-us/web-iq/resources/faq"),
    "S4": ("WebIQ terms", "https://www.microsoft.com/en-us/webiq/webiq-legal"),
    "S5": ("Foundry Browser Automation", "https://learn.microsoft.com/en-us/azure/foundry/agents/how-to/tools/browser-automation"),
    "S6": ("Foundry Web Search", "https://learn.microsoft.com/en-us/azure/foundry/agents/how-to/tools/web-search"),
    "S7": ("Playwright screenshots", "https://playwright.dev/python/docs/screenshots"),
    "S8": ("Playwright CDP connection", "https://playwright.dev/python/docs/api/class-browsertype#browser-type-connect-over-cdp"),
    "S9": ("Chrome extension messaging", "https://developer.chrome.com/docs/extensions/develop/concepts/messaging"),
    "S10": ("Chrome tabs and viewport capture", "https://developer.chrome.com/docs/extensions/reference/api/tabs#method-captureVisibleTab"),
    "P1": ("Authenticated comparison and replay", PORTAL + "/benchmark-report"),
    "P2": ("Implementation and reproducible metadata", r"README.md; api\app\pipeline.py; api\app\bat_research.py; docs\browser-benchmark-runs.json"),
    "P3": ("Architecture reference", "https://github.com/samsonlee0907/anthropic-fsi-skill-demo-with-ghcp"),
}


def color(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=WHITE, line=None, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                   Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if rounded:
        shape.adjustments[0] = 0.04
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size=14, fill=INK, bold=False, font=None, align=None):
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = font or (STRONG if bold else FONT)
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color(fill)
        paragraph.space_after = Pt(6)
        paragraph.line_spacing = 1.05
        if align is not None:
            paragraph.alignment = align
    return shape


def arrow(slide, x, y, w=18, h=12):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(BLUE)
    shape.line.fill.background()
    return shape


def band(slide, value):
    rect(slide, 36, 450, 888, 35, NAVY)
    text(slide, value, 49, 459, 860, 22, 10.5, WHITE, True)


def notes(slide, value, refs=()):
    slide.notes_slide.notes_text_frame.text = (
        value + "\n\nSources and implementation pointers:\n" +
        "\n".join(f"{ref} | {SOURCES[ref][0]}\n{SOURCES[ref][1]}" for ref in refs)
    )


def page(prs, kicker, title, subtitle="", refs=()):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 960, 540)
    rect(slide, 0, 0, 9, 540, BLUE)
    text(slide, kicker.upper(), 36, 20, 850, 20, 10.5, BLUE, True)
    text(slide, title, 36, 42, 888, 38, 24, NAVY, True)
    if subtitle:
        text(slide, subtitle, 36, 86, 888, 33, 13, GRAY)
    text(slide, "Sources: " + ", ".join(refs) if refs else "Prototype discussion | Not product certification",
         36, 492, 835, 20, 8.5, MUTED)
    text(slide, datetime.now(timezone.utc).strftime("%d %b %Y").upper() + " | Product-review research",
         36, 513, 760, 15, 8, MUTED)
    text(slide, f"{len(prs.slides):02d}", 878, 513, 46, 16, 8.5, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, title, body, accent=BLUE, number=None, tag=None, y=124, w=288, h=310):
    rect(slide, x, y, w, h, WHITE, BORDER, rounded=True)
    rect(slide, x, y, w, 65, accent)
    if number:
        text(slide, number, x + 14, y + 11, 44, 44, 31, WHITE, True)
    text(slide, title, x + (62 if number else 14), y + 15,
         w - (76 if number else 28), 46, 16, WHITE, True)
    text(slide, body, x + 15, y + 82, w - 30, h - 132, 12.5)
    if tag:
        rect(slide, x + 11, y + h - 48, w - 22, 37, PALE)
        text(slide, tag, x + 21, y + h - 39, w - 42, 27, 10, accent, True)


def table(slide, headers, rows, widths, y=126, row_h=65, font_size=12):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Pt(36), Pt(y), Pt(sum(widths)),
                                   Pt(40 + row_h * len(rows)))
    grid = shape.table
    for col, width in zip(grid.columns, widths):
        col.width = Pt(width)
    for index, row in enumerate([headers] + rows):
        grid.rows[index].height = Pt(40 if index == 0 else row_h)
        for column, value in enumerate(row):
            cell = grid.cell(index, column)
            cell.text = str(value)
            cell.margin_left = cell.margin_right = Pt(9)
            cell.margin_top = cell.margin_bottom = Pt(7)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = color(NAVY if index == 0 else (WHITE if index % 2 else PALE))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = STRONG if index == 0 or column == 0 else FONT
                paragraph.font.size = Pt(font_size)
                paragraph.font.color.rgb = color(WHITE if index == 0 else INK)
                paragraph.line_spacing = 1.05
    return shape


def code(slide, value, x, y, w, h, size=11.5):
    rect(slide, x, y, w, h, NAVY, rounded=True)
    shape = text(slide, value, x + 15, y + 14, w - 30, h - 28, size, "DCF3FF", font=MONO)
    for paragraph in shape.text_frame.paragraphs:
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.13


def recorded_method_output(prs, row, label, refs):
    """Show public receipt metadata, never expired page prose or private imagery."""
    source = row["inspections"][0]
    capture = next((event for event in row["events"]
                    if event["type"] == "screenshot" and event.get("real_capture")), None)
    fragment = {
        "method": row["method"],
        "candidate_url": source["url"],
        "inspection_status": source["status"],
        "fetched_characters": source.get("fetched_characters", 0),
        "has_browser_capture": bool(source.get("capture")),
    }
    slide = page(prs, f"Method output | {label}", "Read the evidence, not just the success badge",
                 "Actual normalized recording fragment, not the raw vendor response or an invented product finding.", refs)
    code(slide, json.dumps(fragment, indent=2), 36, 137, 555, 226, 10.5)
    if capture:
        code(slide, "Actual listing PNG SHA-256:\n" + capture["sha256"], 36, 375, 555, 60, 9.5)
    else:
        text(slide, "No browser capture is returned by this method.\nMissing or expired prose is not reconstructed here.",
             36, 375, 555, 59, 12, PURPLE, True)
    metrics = [
        (f"{row.get('inspected', 0)} / {len(row['inspections'])} inspected",
         "Discovered review URLs only; the seed is separate."),
        (f"{row.get('real_captures', 0)} actual PNGs",
         "Whole-run count, including the listing. Not a quality score."),
        (f"{row['elapsed_s']:.2f}s | {row['outcome']}",
         "Recorded run time. Completed synthesis is not full verification."),
    ]
    for index, (title, detail) in enumerate(metrics):
        y = 137 + index * 103
        rect(slide, 612, y, 312, 91, PALE, BORDER, rounded=True)
        text(slide, title, 626, y + 12, 284, 28, 17, BLUE, True)
        text(slide, detail, 626, y + 46, 284, 38, 11.5)
    band(slide, "Open the authenticated report for permitted actual captures. Candidate presence is not an endorsement.")
    notes(slide, f"Metadata copied from recording {row['visual_evidence_id']} on {row['server_revision']}. "
          f"The candidate example is {source['url']}. Capture metadata, if shown, is the seed listing's PNG, "
          "not a claim that it depicts the candidate page above. This is a normalized receipt projection, "
          "not a verbatim service response. No publisher content or private PNG is embedded. "
          "Local benchmark execution uses isolated native Chrome via Playwright, not the visitor extension.",
          refs)


def build(output: Path):
    data = json.loads((ROOT / "docs" / "browser-benchmark-runs.json").read_text(encoding="utf-8"))
    fresh = {row["method"]: row for row in data["runs"]
             if row["case_id"] == "sony-wh1000xm6-official"
             and row.get("recording_scope") == "four-stage-site-research"
             and row.get("visual_evidence_id")}
    if any(method not in fresh for method, _ in METHODS):
        raise ValueError("Four genuine, fresh site-research visual recordings are required.")
    if any(len({fresh[method][field] for method, _ in METHODS}) != 1
           for field in ("product_url", "intent", "web_query", "server_revision", "source_commit")):
        raise ValueError("All four comparison inputs and deployed research versions must match.")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    props = prs.core_properties
    props.title = "Product reviews, understood by agents"
    props.subject = "Four site-inspection paths, with a Sony WH-1000XM6 case walkthrough"
    props.author = "Product Due-Diligence Prototype"
    props.created = props.modified = datetime.now(timezone.utc)
    props.keywords = "WebIQ, Playwright, Foundry BAT, site research, evidence"
    revision = fresh["webiq_browse"]["server_revision"]
    receipt_note = (
        f"Fresh matched executions on {revision}; source {fresh['webiq_browse']['source_commit']}. "
        "The numerator counts successfully inspected discovered URLs, not the initial listing. "
        "Runs share URL, intent and primary query; results and access conditions can change. "
        "Counts and elapsed times do not rank answer quality, cost or provider reliability. "
        "No private captures or grounded source prose are reproduced in this presentation."
    )

    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 960, 540, NAVY)
    rect(s, 641, 0, 319, 540, DEEP)
    text(s, "ARCHITECTURE & ENGINEERING DISCUSSION", 58, 68, 562, 36, 13, CYAN, True)
    rect(s, 58, 112, 86, 6, CYAN)
    text(s, "Product reviews,\nunderstood by agents", 58, 139, 563, 135, 36, WHITE, True)
    text(s, "One product URL. A buyer's real priorities.\nA transparent route from sites to a buying brief.",
         58, 289, 554, 86, 16, "C7E1F4")
    text(s, "Sony WH-1000XM6 case walkthrough\nFour methods. Four visible research stages.", 58, 394, 554, 65, 14, WHITE)
    for y, label, abbreviation in ((95, "Understand the listing", "01"), (185, "Find relevant sites", "02"),
                                  (275, "Inspect each page", "03"), (365, "Explain the decision", "04")):
        rect(s, 691, y, 50, 50, BLUE, rounded=True)
        text(s, abbreviation, 695, y + 13, 42, 28, 16, WHITE, True, align=PP_ALIGN.CENTER)
        text(s, label, 758, y + 8, 182, 56, 14, WHITE, True)
    text(s, "Evidence-led prototype | Not product certification or a production guarantee", 58, 493, 565, 30, 9, "C7E1F4")
    notes(s, "Original editable content. The authorized reference informed 960x540 geometry, Segoe UI, navy/cyan colors, "
          "blue rail and card/table layouts only. Its text, logos, pictures and package parts were not copied. "
          "This deck discusses the implemented site-research scope.", ("P1", "P2", "P3"))

    s = page(prs, "01 | Buyer-first framing", "Start with the decision, not a browser demo",
             "Case: Sony WH-1000XM6, official US product page, buyer ceiling below $450.", ("P2",))
    card(s, 36, "The supplied URL", "electronics.sony.com\n/audio/headphones/headband\n/p/wh1000xm6-b\n\nThe URL is a seed, not proof of stock, seller price or variant availability.",
         BLUE, tag="INPUT: ONE PRODUCT PAGE")
    card(s, 336, "What this buyer needs", "10-hour international flights.\nThin prescription glasses.\nWindows laptop + iPhone.\nReliable multipoint switching.\nA durable folding hinge.", PURPLE, tag="INPUT: BRIEF BUYER INTENT")
    card(s, 636, "The useful output", "What is actually observed?\nWhich claims are corroborated?\nWhere are trade-offs or gaps?\nWhat remains unverified?\nWhat should the buyer check next?", TEAL, tag="OUTPUT: ATTRIBUTABLE BRIEF")
    band(s, "The question is not merely whether the headphones are good; it is whether the evidence addresses this buyer.")
    notes(s, "The priorities and seed URL are the actual matched-run inputs, not a new product recommendation. "
          "The price ceiling is the buyer's constraint, not an observed retail price. Model identity is never treated as a verified offer.", ("P2",))

    s = page(prs, "02 | End-to-end route", "Four stages, one inspection choice",
             "A fixed safety scaffold with model-planned queries and evidence-grounded synthesis.", ("P2",))
    stages = [("1", "Listing", "Observe identity,\noffer and claims"),
              ("2", "Discovery", "Plan queries;\nfind review sites"),
              ("3", "Inspection", "Read each eligible\nresult URL"),
              ("4", "Synthesis", "Reconcile claims;\nexplain buyer fit")]
    for index, (num, label, body) in enumerate(stages):
        x = 36 + index * 225
        rect(s, x, 146, 213, 157, PALE, BORDER, rounded=True)
        text(s, num, x + 13, 156, 35, 35, 25, BLUE, True)
        text(s, label, x + 51, 164, 151, 30, 18, NAVY, True)
        text(s, body, x + 14, 222, 185, 66, 14)
        if index < 3:
            arrow(s, x + 215, 212, 10, 12)
    rect(s, 36, 325, 663, 63, NAVY, rounded=True)
    text(s, "Same selected inspector at stages 1 and 3.\nNo hidden switch to a different browsing environment.", 50, 336, 635, 49, 13, WHITE, True)
    text(s, "Approval is separate:\nsimulation only;\ncheckout blocked.", 726, 326, 198, 90, 13, PURPLE, True)
    band(s, "Search excerpts discover candidates. Inspected pages add context. A model must not confuse the two.")
    notes(s, "The runtime builds four research steps. Approval is in post_research_actions, not a fifth research stage. "
          "Dynamic behavior concerns intent/query planning and discovered sources; the orchestrator's safety stages are deterministic. "
          "No unbounded autonomous workflow is claimed.", ("P2",))

    s = page(prs, "03 | Responsibilities", "Separate the services by the question they answer",
             "Different execution and evidence capabilities; not interchangeable names for one browser.", ("S1", "S2", "S5", "S6", "S7", "P2"))
    table(s, ["Component", "Responsibility", "What it does NOT establish"], [
        ["WebIQ Search", "Relevant site URLs and attributed snippets;\nprimary discovery with supplied key.", "That each page was inspected or its\nfacts verified."],
        ["Toolbox web_search", "Alternative discovery without WebIQ key;\nnot a silent fallback after key failure.", "The selected browser environment\nor guaranteed access."],
        ["WebIQ Browse", "Extracted content for a specified URL;\nindex / permitted crawl behavior.", "A viewport screenshot or guaranteed\nlive freshness."],
        ["BAT / Playwright", "Browser observations and genuine captures;\ncloud, local or VM execution.", "Unrestricted access, a clean extraction\nor permission to bypass a site."],
        ["GPT-5.6-Terra", "Intent-aware query planning and\nstructured, cited synthesis.", "Independent truth or a reason to fill\nmissing evidence with invented facts."],
    ], [165, 365, 358], row_h=52, font_size=11.5)
    band(s, "Inspection choice and discovery choice are independent. All four fresh comparisons use WebIQ for discovery.")
    notes(s, "Capabilities paraphrase official sources; enabled behavior is evidenced separately by source and execution metadata. "
          "A supplied invalid WebIQ key produces a visible error, not hidden Toolbox substitution.", ("S1", "S2", "S5", "S6", "S7", "P2"))

    s = page(prs, "04 | Stage 1, Sony baseline", "Establish what the listing actually supports",
             "Identity and offer extraction must distinguish observed fields from unknowns.", ("P2",))
    table(s, ["Baseline field", "Evidence to retain", "Sony-specific check"], [
        ["Identity", "Observed brand, model, SKU and variant.", "WH-1000XM6 vs another generation;\nblack variant vs generic family page."],
        ["Offer", "Observed price, currency, seller and URL.", "Below the buyer's $450 ceiling?\nCurrent price is not inferred from the URL."],
        ["Seller claims", "Testable statements extracted from the page.", "Noise reduction, battery, multipoint,\nfolding design: claims, not verification."],
        ["Access and capture", "Actual inspector, restrictions and capture kind.", "Blocked or empty Sony page means an\nincomplete baseline, even with a screenshot."],
    ], [155, 365, 368], row_h=65)
    band(s, "The product name in the input is not a successful listing inspection. Failures stay visible downstream.")
    notes(s, "This is the extraction checklist, not a claim that every field was returned in each live run. "
          "The case-level receipts and private replay distinguish successful observations, unknowns and restrictions. "
          "Do not quote a current price unless it was actually observed.", ("P2",))

    s = page(prs, "05 | Stage 2, intent into queries", "Translate buyer priorities into research questions",
             "Illustrative planning logic. The matched benchmark pins a shared primary query.", ("S2", "S6", "P2"))
    table(s, ["Buyer priority", "Illustrative query focus", "Useful source type"], [
        ["Glasses + long flights", "Comfort, seal and ANC with glasses.", "Independent acoustic lab + owner accounts."],
        ["Laptop + iPhone", "Multipoint switching, call hand-off, firmware.", "Repeated owner reports + technical testing."],
        ["Durable folding hinge", "Long-term failures, repair and warranty.", "Owner communities + repair documentation."],
        ["Battery + budget", "Real test conditions, current seller offer.", "Measured review + observed product listing."],
    ], [185, 335, 368], row_h=61, font_size=12)
    text(s, "Matched primary: Sony WH-1000XM6 Headphones independent owner review long term problems",
         36, 418, 888, 29, 11, BLUE, True)
    band(s, "Discover platforms from relevance. Reddit is one possible destination, never the research template.")
    notes(s, "Illustrative questions are not represented as the exact model output. Fresh four-path runs use the same explicit primary "
          "query override and the application's supplementary search. Returned URL sets can differ over time. "
          "This distinction prevents a source-list difference from being misattributed solely to the inspector.", ("S2", "S6", "P2"))

    s = page(prs, "06 | Discovery request boundary", "Search finds candidates; it does not inspect them",
             "Prototype-shaped code with placeholders, not an embedded credential or a vendor code sample.", ("S2", "S6", "P2"))
    code(s, 'await client.post(\n'
         '    "https://api.microsoft.ai/v3/search/web",\n'
         '    headers={"x-apikey": request_key},\n'
         '    json={"query": primary_query,\n'
         '          "maxResults": 8,\n'
         '          "contentFormat": "passage"},\n'
         ')\n\n'
         '# Normalize URLs, classify source type,\n'
         '# keep citations, deduplicate candidates.\n'
         '# No key: use Toolbox web_search instead.',
         36, 132, 542, 285, 11.5)
    text(s, "What the agent receives", 603, 140, 315, 35, 18, NAVY, True)
    text(s, "URL and title.\nAttributed short excerpt.\nSource category and provenance.\n\nNot yet a viewed page.\nNot an authenticity certificate.",
         603, 192, 315, 214, 14)
    band(s, "Key presence chooses the discovery route. A failing selected provider is an explicit gap, not a fake success.")
    notes(s, "The snippet illustrates the implemented endpoint and boundary. See api/app/webiq.py for full validation, request settings "
          "and error handling. Excerpt display is capped at 200 characters or a lower publisher limit, with direct attribution. "
          "That cap alone is not a compliance certification.", ("S2", "S3", "S4", "S6", "P2"))

    s = page(prs, "07 | Stage 3, the routing invariant", "Inspect every unique eligible result the same way",
             "Equivalent control-flow illustration; actual implementation also emits progress and enforces deadlines.", ("P2",))
    code(s, 'inspector = selected_method\n'
         'baseline = await digest(product_url, inspector)\n'
         'sources = await discover(baseline, user_intent)\n\n'
         'for source in unique_eligible_urls(sources):\n'
         '    observation = await inspect(source.url,\n'
         '                                inspector)\n'
         '    emit_source_outcome(observation)\n\n'
         'brief = await synthesize(baseline, sources,\n'
         '                         observations)',
         36, 133, 542, 286, 12)
    text(s, "Never hide these outcomes", 603, 140, 315, 35, 18, NAVY, True)
    text(s, "Inspected with page context.\nRestricted by policy or challenge.\nPending crawl or empty content.\nFailed or timed out.\n\nSearch excerpt != page observation.",
         603, 192, 315, 214, 14)
    band(s, "No fixed two-site script. No split where some review pages silently switch to another inspector.")
    notes(s, "The selected inspector digests the seed and loops over unique eligible discovered sources. "
          "Policy skips are recorded without evasion. The 20-minute total run budget and per-call deadlines may yield partial coverage. "
          "Every attempted or skipped outcome remains diagnosable.", ("P2",))

    s = page(prs, "08 | WebIQ Browse path", "Extract page content without owning a browser session",
             "Suitable when attributed text is sufficient and indexed/crawl behavior is acceptable.", ("S1", "S3", "P2"))
    code(s, 'await client.post(\n'
         '    "https://api.microsoft.ai/v3/browse",\n'
         '    headers={"x-apikey": request_key},\n'
         '    json={\n'
         '        "url": review_url,\n'
         '        "liveCrawl": "fallback",\n'
         '        "renderDynamicPages": True,\n'
         '        "contentFormat": "markdown",\n'
         '        "maxLength": 15000,\n'
         '    },\n'
         ')', 36, 132, 542, 286, 12)
    text(s, "Capabilities and limits", 603, 140, 315, 35, 18, NAVY, True)
    text(s, "Content plus freshness metadata.\nFallback crawl may be pending.\nIndex presence affects availability.\nDynamic rendering is not a viewport.\n\nNo native PNG screenshot here.",
         603, 192, 315, 214, 14)
    band(s, "Use freshness fields and response status; do not label indexed content as a guaranteed current storefront.")
    notes(s, "WebIQ supports none/fallback/force crawl modes; this prototype selects fallback and permits dynamic rendering. "
          "renderDynamicPages is supported with fallback. The requested extraction length is bounded. "
          "A diagnostic card is not a live browser capture and is excluded from genuine screenshot counts.", ("S1", "S3", "P2"))

    s = page(prs, "Method specification | WebIQ Browse", "A URL-to-content API, not remote desktop control",
             "Distinguish the published service contract from the settings selected by this prototype.", ("S1", "P2"))
    table(s, ["Technical aspect", "Service contract", "Prototype selection"], [
        ["Transport and identity", "HTTPS POST /v3/browse;\nAPI key or Microsoft Entra authentication.", "Request-scoped x-apikey;\nno key saved with the report."],
        ["Content format and size", "text / html / markdown;\nmaxLength default 10,000, maximum 500,000.", "Markdown; requested maxLength 15,000.\nThe returned content can be shorter."],
        ["Crawl behavior", "none / fallback / force;\ndynamic rendering supported with fallback.", "fallback + renderDynamicPages=true;\n202 stays pending, not inspected."],
        ["Freshness and status", "url, title, content, traceId;\ncrawledAt and lastUpdatedAt when returned.", "Read returned metadata and actual status.\nNo freshness claim from a missing field."],
    ], [176, 356, 356], row_h=65, font_size=11.5)
    band(s, "liveCrawlSucceeded is not a true-valued success flag: docs specify false for stale force results; otherwise omitted.")
    notes(s, "Official Browse reference verified for this expansion. A fallback crawl is not equivalent to forcing every page fresh. "
          "HTTP 200 content, 202 pending and explicit restriction/rate-limit errors remain distinct. "
          "The schema fields listed here are documentation, not an assertion that every receipt includes them.", ("S1", "P2"))
    recorded_method_output(prs, fresh["webiq_browse"], "WebIQ Browse", ("S1", "P1", "P2"))

    s = page(prs, "09 | Azure BAT Tool path", "A governed cloud browser with observable page state",
             "Foundry Toolbox creates the browser session; Playwright reads through CDP.", ("S5", "S7", "P2"))
    boxes = [("Portal identity", "Pinned local identity or\nhosted managed identity"),
             ("Foundry Toolbox", "browser_automation\ncreate_session"),
             ("Cloud browser", "Azure Playwright workspace\nCDP endpoint"),
             ("Observation", "DOM-derived fields and\nactual viewport PNG")]
    for index, (title, body) in enumerate(boxes):
        x = 36 + index * 225
        rect(s, x, 150, 213, 155, PALE, BORDER, rounded=True)
        text(s, title, x + 12, 168, 189, 44, 17, NAVY, True)
        text(s, body, x + 12, 235, 189, 65, 12)
        if index < 3:
            arrow(s, x + 215, 219, 10, 12)
    text(s, "A cloud browser is not the user's local session. Challenges, login, page design and policy can still block evidence.",
         36, 338, 888, 60, 15, PURPLE, True)
    band(s, "A screenshot proves what that browser displayed; it does not prove a clean extraction or a correct seller claim.")
    notes(s, "The prototype uses the existing bat-toolbox and Playwright connection. Browser sessions are cleaned up. "
          "The selected BAT method remains selected for discovered review pages. The screenshot path returns genuine PNG bytes, "
          "not a generated product visualization.", ("S5", "S7", "P2"))

    s = page(prs, "Method specification | Azure BAT Tool", "Govern the session, then observe through Playwright",
             "Browser Automation is generally available; broader vendor capabilities are not all enabled here.", ("S5", "S7", "S8", "P2"))
    table(s, ["Layer", "What actually runs", "Technical boundary"], [
        ["Tool orchestration", "Foundry Toolbox browser_automation\ncreate_session tool and connection.", "Toolbox governs connection/auth/policy.\nNo arbitrary CDP endpoint from the buyer."],
        ["Browser service", "Azure Playwright workspace;\nreturned secure WebSocket CDP endpoint.", "Cloud execution is a distinct environment,\nnot the visitor's browser session."],
        ["Observation", "Playwright attaches to Chromium via CDP;\nreads DOM / JSON-LD and viewport PNG.", "CDP is Chromium-only and lower fidelity\nthan the native Playwright protocol."],
        ["Enabled vs available", "Read-only listing and review inspection;\nbounded waits and session cleanup.", "Live view, take-control and broader\ninteraction are not portal features here."],
    ], [176, 356, 356], row_h=65, font_size=11.5)
    band(s, "A protected or challenge page can produce a real PNG with zero useful product text. Preserve both outcomes.")
    notes(s, "Official Foundry documentation describes a generally available MCP browser tool backed by Playwright Workspaces. "
          "The prototype asks Toolbox for its browser session and connects to the returned CDP endpoint. "
          "This slide separates vendor surface area from implemented behavior, not a claim of end-to-end feature parity.", ("S5", "S7", "S8", "P2"))
    recorded_method_output(prs, fresh["live"], "Azure BAT Tool", ("S5", "S7", "P1", "P2"))

    s = page(prs, "10 | Local and VM Playwright paths", "Same observation protocol, different browser placement",
             "Choose the environment deliberately; it affects session context and operational dependencies.", ("S7", "P2"))
    table(s, ["Aspect", "Local Playwright", "Azure VM + Playwright"], [
        ["Execution", "Chrome / Edge extension and local\nobservation bridge.", "Existing Windows VM and named\nPlaywright worker task."],
        ["Control path", "Authenticated research stream sends\ncommands; browser returns observations.", "Authenticated heartbeat, command queue\nand observation responses."],
        ["Evidence", "Actual local page text and viewport PNG;\nbenchmark uses isolated native Chrome.", "Actual VM page text and viewport PNG;\nnot borrowed from the user's desktop."],
        ["Failure mode", "Extension unavailable, tab/session issues,\nsite challenge or extraction failure.", "Worker offline, stale credentials, VM\navailability or site challenge."],
    ], [140, 374, 374], row_h=64, font_size=12)
    band(s, "A VM being powered on is not sufficient. Verify the worker's fresh authenticated heartbeat.")
    notes(s, "The live benchmark runs the same Local Playwright observation protocol using native Chrome in an isolated runner profile. "
          "It does not borrow the VM or relabel cloud BAT as local execution. The actual worker VM is vm-pdd-win-chrome. "
          "Passcode rotation also requires the worker launcher to be updated; no credentials are included here.", ("S7", "P2"))

    s = page(prs, "Method identity | Local Playwright", "One local observation contract, two implementations",
             "The existing portal label is retained. Its interactive bridge is a Chrome extension, not a Playwright client.", ("S7", "S9", "S10", "P2"))
    card(s, 36, "Interactive portal path",
         "Portal receives a research command.\nContent script relays a message.\nExtension navigates a local tab.\nChrome scripting reads the page.\ncaptureVisibleTab returns a PNG.",
         BLUE, tag="BROWSER_EXTENSION IMPLEMENTATION", w=432)
    card(s, 492, "Matched benchmark runner",
         "An isolated native Chrome profile.\nPlaywright drives the local page.\nThe same observation fields return.\nNo extension session is borrowed.\nThis is not an extension E2E test.",
         PURPLE, tag="ISOLATED PLAYWRIGHT IMPLEMENTATION", w=432)
    band(s, "Same normalized evidence contract does not mean identical permissions, session state, viewport or access behavior.")
    notes(s, "See browser_extension/background.js, portal_content.js and scripts/refresh_benchmark_report.py. "
          "The Local Playwright label historically names an execution option, not its literal interactive library. "
          "Benchmark timings must not be claimed as direct measurements of the visitor's Chrome extension.", ("S7", "S9", "S10", "P2"))

    s = page(prs, "Method specification | Local browser", "Understand the local bridge's permissions and limits",
             "Installed and connected by the visitor; it never silently becomes cloud BAT or the VM.", ("S9", "S10", "P2"))
    table(s, ["Technical aspect", "Interactive extension behavior", "Interpretation"], [
        ["Runtime", "Manifest V3 service worker;\nportal content-script message bridge.", "Requires an installed extension and\nan available compatible browser."],
        ["Permissions", "tabs, activeTab, scripting;\nHTTPS host access plus loopback hosts.", "A broad local capability surface;\nreview before installing."],
        ["Page output", "DOM / JSON-LD-derived observation;\nChrome tabs.captureVisibleTab PNG.", "A viewport capture, not an entire\nsite archive or independently verified fact."],
        ["Status reporting", "Bridge sets status_code=200 after a\nloaded tab, subject to extraction checks.", "Client-reported load status is NOT a\nmeasured origin HTTP response code."],
    ], [176, 356, 356], row_h=65, font_size=11.5)
    band(s, "The local browser's context and network are different, not a guarantee of access or permission to evade a policy.")
    notes(s, "Permissions are taken from the actual manifest; the HTTP status caveat comes from the actual background script. "
          "Chrome's capture API works on the visible area of the active tab. Browser restrictions and tab state still apply. "
          "The local benchmark output on the next slide comes from the isolated Playwright runner, not this extension.", ("S9", "S10", "P2"))
    recorded_method_output(prs, fresh["client_browser"], "Local Playwright", ("S7", "S9", "P1", "P2"))

    s = page(prs, "Method specification | Azure VM + Playwright", "A self-operated browser worker, not another managed tool",
             "The portal queues work; the existing Windows worker executes it and returns observations.", ("S7", "S8", "P2"))
    table(s, ["Technical aspect", "Configured worker implementation", "Operational implication"], [
        ["Execution and identity", "vm-pdd-win-chrome; named scheduled task;\nauthenticated polling and heartbeat.", "Portal worker status must be fresh.\nVM power state alone is insufficient."],
        ["Browser and profile", "Native Chrome + persistent worker profile;\n1366 x 900 viewport, en-US locale.", "Separate from the buyer's browser.\nNo guarantee of equivalent sessions."],
        ["Display mode", "Modern native --headless=new by default;\noptional headed mode or CDP attachment.", "Native Chrome does not mean a\nvisible desktop window is in use."],
        ["Navigation and capture", "domcontentloaded; 26-second goto budget;\nactual response status where available.", "Bounded subsequent waits;\nPNG screenshot(full_page=False)."],
    ], [176, 356, 356], row_h=65, font_size=11.5)
    band(s, "The operator owns VM availability, browser/profile hygiene, patching, worker credentials and ongoing compute cost.")
    notes(s, "Values are implementation specifications from scripts/windows_playwright_worker.py, not default values claimed "
          "for every Playwright application or verified measurements of all historical executions. "
          "The worker supports an existing CDP connection as an alternative to launching its persistent context. "
          "Private credentials, browser content and machine session data are excluded.", ("S7", "S8", "P2"))
    recorded_method_output(prs, fresh["vm_playwright"], "Azure VM + Playwright", ("S7", "P1", "P2"))

    s = page(prs, "11 | Sony, discovered sites", "Look at the actual candidate mix before drawing conclusions",
             "Fresh execution metadata only. A domain appearing in search is not an endorsement.", ("P1", "P2"))
    domains = []
    for method, _ in METHODS:
        for row in fresh[method].get("inspections", []):
            domain = urlsplit(row["url"]).hostname or ""
            if domain and domain not in domains:
                domains.append(domain)
    rows = [[domain, " | ".join(str(sum((urlsplit(row["url"]).hostname or "") == domain
                                     for row in fresh[method].get("inspections", [])))
                               for method, _ in METHODS)]
            for domain in domains[:5]]
    table(s, ["Observed domain", "Recorded URL outcomes: WebIQ | BAT | Local | VM"], rows,
          [420, 468], row_h=49, font_size=12)
    text(s, "Counts include restricted / failed URLs. This is coverage metadata, not source credibility or corroboration.",
         36, 421, 888, 27, 11, BLUE, True)
    band(s, "Use labs for measurements, owners for recurring experience, and independent evidence to challenge seller claims.")
    notes(s, receipt_note + " Domains are selected in first-observed order from the fresh runs. Source categories require content review; "
          "the deck does not automatically certify a site's editorial independence or expertise.", ("P1", "P2"))

    s = page(prs, "12 | Sony, source-level outcomes", "The same URL can produce different inspection evidence",
             "Selected overlapping URLs from fresh recordings; absent means not returned in that run.", ("P1", "P2"))
    candidates = list(dict.fromkeys(row["url"] for method, _ in METHODS
                                   for row in fresh[method].get("inspections", [])))
    candidates.sort(key=lambda url: -sum(any(row["url"] == url for row in fresh[method].get("inspections", []))
                                        for method, _ in METHODS))
    rows = []
    for url in candidates[:4]:
        values = []
        for method, _ in METHODS:
            row = next((row for row in fresh[method].get("inspections", []) if row["url"] == url), None)
            values.append(f"{row['status']}\n{row.get('fetched_characters', 0):,} chars" if row else "Not returned")
        label = (urlsplit(url).hostname or "") + "\n" + urlsplit(url).path[-48:]
        rows.append([label] + values)
    table(s, ["Discovered URL", "WebIQ", "BAT", "Local", "VM"], rows,
          [320, 142, 142, 142, 142], row_h=65, font_size=10.5)
    band(s, "A failed extraction is an evidence gap. Do not infer a provider defect or access-policy root cause from one run.")
    notes(s, receipt_note + "\nFull URLs for the table:\n" + "\n".join(candidates[:4]), ("P1", "P2"))

    s = page(prs, "13 | Fresh four-path receipts", "Compare what was actually completed",
             "Matched Sony seed, buyer intent, primary query and deployed research source.", ("P1", "P2"))
    rows = []
    for method, label in METHODS:
        row = fresh[method]
        rows.append([label, row["outcome"], f"{row.get('inspected', 0)} / {len(row.get('inspections', []))}",
                     str(row.get("real_captures", 0)), f"{row['elapsed_s']:.1f}s",
                     row.get("synthesis_status", "Not returned")])
    table(s, ["Method", "Outcome", "Inspected /\nURL outcomes", "Actual\nPNGs", "Elapsed", "GPT synthesis"],
          rows, [206, 103, 145, 97, 105, 232], row_h=64, font_size=11.5)
    text(s, f"Research revision: {revision} | Prior attempts stay in the report; this table uses fresh four-stage runs.",
         36, 426, 888, 24, 10, BLUE, True)
    band(s, "Coverage and capture counts do not score answer quality. Zero WebIQ viewport captures is an intentional boundary.")
    notes(s, receipt_note, ("P1", "P2"))

    s = page(prs, "14 | Capture truth", "Three evidence states the UI must not blur",
             "A screenshot, extracted prose and missing evidence answer different diagnostic questions.", ("S1", "S7", "P1"))
    card(s, 36, "Extracted page text", "WebIQ Browse returns content.\nFreshness fields matter.\nA diagnostic rendering is not a captured website.\nNo synthetic PNG is counted as a browser screenshot.",
         BLUE, tag="OBSERVED TEXT != VIEWPORT")
    card(s, 336, "Actual browser capture", "A browser method retains real PNG bytes.\nSHA-256 binds the capture.\nA challenge screen is still a genuine image, but not a successful product extraction.",
         PURPLE, tag="CAPTURE != CORROBORATION")
    card(s, 636, "Unavailable evidence", "A restriction or timeout stays visible.\nExpired context stays expired.\nA replay never reconstructs evidence or silently reruns research to appear complete.",
         TEAL, tag="UNKNOWN != VERIFIED")
    band(s, "Show the provenance and access outcome next to the content, not hidden only in raw diagnostics.")
    notes(s, "Only genuine browser captures are counted and packaged as verified PNGs. Private captures remain in the authenticated deployment, "
          "not the shared presentation. WebIQ-grounded prose is transient, and absence of a current replay payload is explicit.", ("S1", "S7", "P1"))

    s = page(prs, "15 | Stage 4, claim versus reality", "Reason against buyer priorities, not generic sentiment",
             "Illustrative evaluation questions, not asserted Sony product findings.", ("P2",))
    table(s, ["Claim or priority", "Evidence needed", "Safe decision language"], [
        ["Comfort with glasses", "Test conditions, seal effects, wearing duration;\nowner context rather than one star score.", "Observed under these conditions;\nthis buyer's fit still needs a trial."],
        ["Reliable multipoint", "Exact devices, OS, firmware and switching test.", "Corroborated only for the tested setup;\nother combinations remain uncertain."],
        ["Durable folding hinge", "Independent long-term reports, denominators,\nrepair data and warranty terms.", "Reported incidents are a signal,\nnot a measured failure rate."],
        ["Fits budget", "Current observed offer, seller and variant.", "Price was / was not observed;\ndo not invent a qualifying offer."],
    ], [187, 387, 314], row_h=64, font_size=11.5)
    band(s, "Cite a concrete source for each claim check. Distinguish measurement, anecdote, seller statement and inference.")
    notes(s, "These are analytical questions designed from the actual buyer intent. No unlicensed publisher prose or private model output "
          "is reproduced. The runtime returns cited claim rows, user-intent impact and evidence gaps; review-quality signals are heuristics, "
          "not a guaranteed authenticity classification.", ("P2",))

    s = page(prs, "16 | Synthesis contract", "Use structured output with explicit evidence constraints",
             "Implementation excerpt, shortened to show the contract rather than every prompt rule.", ("P2",))
    code(s, 'synth_llm = model.with_structured_output(\n'
         '    DueDiligenceSynthesisSchema\n'
         ')\n'
         'synthesized = await bounded_await(\n'
         '    synth_llm.ainvoke(prompt),\n'
         '    32.0,\n'
         '    "GPT due-diligence synthesis",\n'
         ')\n\n'
         '# Inputs: baseline, intent, cited search,\n'
         '# actual inspected-page context + gaps.',
         36, 132, 542, 286, 12)
    text(s, "Critical prompt rules", 603, 140, 315, 35, 18, NAVY, True)
    text(s, "Never invent facts or measurements.\nEvery claim row needs a URL.\nExplain impact on buyer priorities.\nList missing / restricted evidence.\nTreat external text as untrusted.\nDo not promote snippets to visits.",
         603, 192, 315, 214, 13.5)
    band(s, "GPT-5.6-Terra remains the research model. A schema shapes output; it does not independently verify the answer.")
    notes(s, "See synthesize_due_diligence_with_gpt in api/app/bat_research.py. The structured schema contains claim rows, "
          "summary, recommendation and evidence gaps. This shortened excerpt retains the real timeout. "
          "Model errors remain explicit partial evidence; sample content never replaces failed inference.", ("P2",))

    s = page(prs, "17 | Evidence UX and replay", "Make the research understandable without reading a log",
             "Editable schematic, not a fabricated portal screenshot.", ("P1", "P2"))
    for index, (title, detail) in enumerate((
            ("Listing", "Observed identity, offer and seller claims"),
            ("Discovery", "Source links and candidate coverage"),
            ("Inspection", "Per-URL context, restriction and capture"),
            ("Synthesis", "Buying brief, citations and evidence gaps"))):
        y = 139 + index * 66
        rect(s, 36, y, 532, 56, PALE, BORDER, rounded=True)
        text(s, str(index + 1), 49, y + 13, 28, 30, 18, BLUE, True)
        text(s, title, 90, y + 7, 125, 35, 14, NAVY, True)
        text(s, detail, 222, y + 10, 331, 40, 11.5)
    text(s, "Replay boundaries", 601, 141, 323, 40, 19, PURPLE, True)
    text(s, "Read recordings only.\nKeep earlier attempts selectable.\nExpand real page observations.\nLoad genuine PNG galleries.\nLabel expired context honestly.\nKeep raw events collapsed.",
         601, 199, 323, 215, 14)
    band(s, "Event playback never calls the research endpoint or invokes a new model.")
    notes(s, "The report provides four expandable visual steps and separate event timing controls. "
          "Recordings without retained observations explicitly say metadata-only. Authenticated GETs load replay evidence and captures. "
          "Historical complete-run timings are not relabelled as fresh four-stage measurements.", ("P1", "P2"))

    s = page(prs, "18 | Efficiency and method selection", "Choose the lightest path that can answer the question",
             "Efficiency means relevant evidence and explicit bounds, not hiding missing work.", ("S1", "S5", "S7", "P2"))
    table(s, ["Need", "Useful starting choice", "Trade-off to inspect"], [
        ["Attributed textual research", "WebIQ Browse + WebIQ Search.", "Indexed freshness, content coverage,\npending crawls and no viewport."],
        ["Cloud browser observations", "Azure BAT Tool + available discovery.", "Session lifecycle, site restrictions,\ncloud access conditions."],
        ["Local browser context", "Local Playwright + available discovery.", "Extension readiness, tab context,\nprivacy and execution ownership."],
        ["Separately operated browser", "Azure VM + Playwright + discovery.", "VM cost, worker heartbeat, credentials,\npatching and operational maintenance."],
    ], [228, 324, 336], row_h=63, font_size=11.5)
    band(s, "No measured cost or answer-quality ranking is claimed. The report shows observed execution, not a universal winner.")
    notes(s, "The prototype deduplicates URLs, bounds each call, cleans up browser sessions and has a 1,200-second total research budget. "
          "It does not optimize by silently dropping a fixed number of sources. Per-query cost and recommendation accuracy need a separate, "
          "rights-aware evaluation design.", ("S1", "S5", "S7", "P2"))

    s = page(prs, "19 | Operating boundaries", "Keep access, retention and human actions explicit",
             "Implemented prototype controls are not production identity, legal approval or an SLA.", ("S3", "S4", "S6", "P2"))
    table(s, ["Boundary", "Implemented now", "Production discussion"], [
        ["Access", "HTTPS passcode session, HttpOnly cookie,\nCSRF checks and protected replay.", "Per-user identity, roles, ownership,\nquotas and audit requirements."],
        ["Secrets", "Request-scoped WebIQ key; server identity;\nworker credential updated on rotation.", "Secret lifecycle, monitoring and\nleast-privilege access review."],
        ["Evidence lifecycle", "Grounded text and synthesis: <=15-min RAM.\nPrivate permitted browser PNGs are separate.", "Approved durable storage, lifecycle,\nprovider terms and deletion obligations."],
        ["Human action", "Optional approval ticket is simulation only.\nNo retailer cart mutation or checkout.", "Commerce integration requires a new\nexplicit authorization and risk design."],
    ], [155, 380, 353], row_h=64, font_size=11.5)
    band(s, "A public page or a successful fetch is not a blanket reuse license. Keep citations and review provider terms.")
    notes(s, "Single replica/worker preserves bounded in-memory context but restart discards it. Private replay is image-packaged, "
          "not governed durable cloud storage. The presentation includes no keys, private captures or grounded publisher prose. "
          "The cap on displayed excerpts is a technical control, not a legal compliance opinion.", ("S3", "S4", "S6", "P2"))

    s = page(prs, "20 | Discussion", "Agree an evidence-quality bar before adding autonomy",
             "The Sony example makes both the value and the remaining uncertainty inspectable.", ("P1", "P2"))
    card(s, 36, "Evidence quality", "Which priorities need measurements?\nWhich can use owner experience?\nWhat counts as independent evidence?\nHow should conflicting sources be handled?",
         BLUE, tag="DECISION: ACCEPTABLE EVIDENCE")
    card(s, 336, "Execution placement", "Do we need rendered page state?\nIs attributed text enough?\nWho operates local / VM browsers?\nWhich budget and timeout is acceptable?",
         PURPLE, tag="DECISION: INSPECTION ENVIRONMENT")
    card(s, 636, "Human accountability", "Who checks the final recommendation?\nWhat must remain unknown?\nWhich evidence can be retained?\nWhat actions must stay simulation-only?",
         TEAL, tag="DECISION: GOVERNANCE")
    band(s, "Success is an explainable, source-backed decision with visible gaps, not simply a green completed badge.")
    notes(s, "This discussion is about engineering and evidence quality, not a product endorsement. Fresh executions can be partial "
          "even when the synthesis stage returns output. Review the authenticated report for actual source outcomes and genuine captures.", ("P1", "P2"))

    source_keys = list(SOURCES)
    for start in range(0, len(source_keys), 8):
        refs = tuple(source_keys[start:start + 8])
        s = page(prs, f"Appendix {start // 8 + 1} | Official documentation",
                 "Sources, implementation and reproducible receipts",
                 "Clickable links and slide-specific context are included in speaker notes.", refs)
        for index, key in enumerate(refs):
            column, row = divmod(index, 4)
            x, y = 36 + column * 450, 137 + row * 70
            rect(s, x, y, 438, 58, PALE if row % 2 == 0 else WHITE)
            text(s, key, x + 9, y + 16, 35, 25, 12, BLUE, True)
            link = text(s, SOURCES[key][0], x + 50, y + 15, 375, 39, 12, NAVY)
            if SOURCES[key][1].startswith("https://"):
                link.text_frame.paragraphs[0].runs[0].hyperlink.address = SOURCES[key][1]
        band(s, "Vendor capabilities can change. Source code and execution receipts establish what this prototype actually used.")
        notes(s, "Official owner documentation was prioritized. Source references and timestamps are not proof that every documented feature "
              "is enabled. Rebuild with scripts/build_discussion_deck.py; only public metadata is read. No cloud calls or inference occur.", refs)

    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + Pt(1)
                    or shape.top + shape.height > prs.slide_height + Pt(1)):
                raise ValueError(f"Slide {index} has an out-of-bounds shape: {shape.name}")
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    check = Presentation(output)
    if len(check.slides) != 32 or any(not slide.has_notes_slide for slide in check.slides):
        raise ValueError("Expected thirty-two slides with notes and citations.")
    print(json.dumps({"path": str(output.resolve()), "slides": len(check.slides),
                      "native_editable_shapes": sum(len(slide.shapes) for slide in check.slides),
                      "research_revision": revision, "built_at": datetime.now(timezone.utc).isoformat()}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "docs" / "Product-Review-Agent-Discussion-Expanded.pptx")
    build(parser.parse_args().output)
