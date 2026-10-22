"""Create editable solution diagrams with official Microsoft Azure icons.

The user-supplied WorkIQ briefing is a visual reference only. Its services,
identities, tenant details and multi-user guarantees are not imported.
There are no research, deployment or provider calls.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_educational_deck import (
    AMBER, BLUE, BORDER, GRAY, INK, NAVY, PALE, PURPLE, TEAL, WHITE,
    rect, render_and_audit, rgb, text,
)

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Product-Review-Agent-Solution-Concept.pptx"
ICON_DIR = ROOT / "docs" / "assets" / "azure-icons"
ICON_KEYS = (
    "container-apps-environment", "foundry-models", "storage-account",
    "virtual-machine", "managed-identity", "toolbox", "browser",
)
DIAGRAMS = (
    "01-architecture", "02-research-flow", "03-request-sequence",
    "04-routing-choices", "05-evidence-boundary",
)
REFERENCE = {
    "filename": "Foundry-WorkIQ-Tool-Blob-Security-Briefing.pptx",
    "sha256": "9afe56b777a27f269bbe350ccc4dfaf009f1031c3ab2a8ecbf6fb0d6641d98d9",
    "usage": "Visual conventions only; no source slides, images or tenant content copied.",
}
SOURCES = {
    "Toolbox": "https://learn.microsoft.com/en-us/azure/foundry/agents/how-to/tools/toolbox",
    "Blob identity": "https://learn.microsoft.com/en-us/azure/storage/blobs/authorize-access-azure-active-directory",
    "Drawing API": "https://python-pptx.readthedocs.io/en/latest/user/autoshapes.html",
    "Azure icons": "https://learn.microsoft.com/en-us/azure/architecture/icons/",
}
IMPLEMENTATION = (
    "README.md; docs/DEPLOYMENT.md; api/app/pipeline.py; api/app/bat_research.py; "
    "api/app/community_exploration.py; api/app/evidence_archive.py; api/app/auth.py"
)


def validate_icon_assets():
    manifest = json.loads((ICON_DIR / "manifest.json").read_text(encoding="utf-8"))
    if manifest["source_page"] != SOURCES["Azure icons"] or manifest["svg_content_changed"] is not False:
        raise ValueError("Use the attributed, unchanged official Microsoft icon assets.")
    for key in ICON_KEYS:
        for extension in ("svg", "png"):
            content = (ICON_DIR / f"{key}.{extension}").read_bytes()
            if hashlib.sha256(content).hexdigest() != manifest["icons"][key][extension + "_sha256"]:
                raise ValueError(f"Azure icon hash mismatch: {key}.{extension}")
    if hashlib.sha256((ICON_DIR / "Microsoft_Terms_of_Use.pdf").read_bytes()).hexdigest() != manifest["terms_sha256"]:
        raise ValueError("Microsoft icon terms do not match their source receipt.")
    return manifest


def service_icon(slide, key, x, y, size=28):
    if key not in ICON_KEYS:
        raise ValueError(f"Unknown official icon: {key}")
    # Width only preserves the artwork's aspect ratio, including transparent padding.
    shape = slide.shapes.add_picture(str(ICON_DIR / f"{key}.png"), Pt(x), Pt(y), width=Pt(size))
    shape.name = "azure-icon-" + key
    shape._element.nvPicPr.cNvPr.set("descr", "Official Microsoft icon: " + key)
    return shape


def node(slide, title, body, x, y, width, height, color=BLUE, size=11.5, icon_key=None):
    shape = rect(slide, x, y, width, height, WHITE, color)
    shape.name = "concept-node"
    inset = 46 if icon_key else 10
    if icon_key:
        service_icon(slide, icon_key, x + 10, y + 10)
    text(slide, title, x + inset, y + 9, width - inset - 10, 35, 13, color, True)
    if body:
        text(slide, body, x + 10, y + 41, width - 20, height - 47, size, INK)
    return shape


def zone(slide, label, x, y, width, height, color=BLUE, icon_key=None):
    shape = rect(slide, x, y, width, height, PALE, color)
    shape.name = "responsibility-group"
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    inset = 39 if icon_key else 9
    if icon_key:
        service_icon(slide, icon_key, x + 9, y + 4, 24)
    text(slide, label, x + inset, y + 6, width - inset - 9, 25, 9.3, color, True)
    return shape


def route(slide, points, color=BLUE, dashed=False, arrow=True):
    for index, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        shape.name = "concept-flow"
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(1.7)
        if dashed:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if arrow and index == len(points) - 2:
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            shape._element.spPr.get_or_add_ln().append(end)


def number(slide, value, x, y, color=BLUE):
    shape = rect(slide, x - 9, y - 9, 18, 18, color, kind=MSO_SHAPE.OVAL)
    shape.name = "flow-number"
    text(slide, str(value), x - 8, y - 6.4, 16, 14, 9, WHITE, True, align=PP_ALIGN.CENTER)


def band(slide, value, icon_key=None):
    rect(slide, 36, 467, 888, 34, NAVY, kind=MSO_SHAPE.RECTANGLE)
    if icon_key:
        rect(slide, 46, 471, 26, 26, WHITE)
        service_icon(slide, icon_key, 49, 474, 20)
    inset = 82 if icon_key else 49
    text(slide, value, inset, 475, 911 - inset, 23, 10.5, WHITE, align=PP_ALIGN.CENTER)


def page(prs, label, title, subtitle, narrative, refs=()):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 960, 540, WHITE, kind=MSO_SHAPE.RECTANGLE)
    rect(slide, 0, 0, 9, 540, BLUE, kind=MSO_SHAPE.RECTANGLE)
    text(slide, label.upper(), 36, 20, 888, 18, 10.5, BLUE, True)
    text(slide, title, 36, 44, 888, 38, 25, NAVY, True)
    text(slide, subtitle, 36, 88, 888, 34, 12.5, GRAY)
    text(slide, "Solution concept | Current prototype | 21 October 2026", 36, 519, 760, 13, 8, GRAY)
    text(slide, f"{len(prs.slides):02}", 880, 517, 44, 17, 9, GRAY, align=PP_ALIGN.RIGHT)
    for index, key in enumerate(dict.fromkeys((*refs, "Azure icons"))):
        shape = text(slide, key, 36 + index * 180, 505, 170, 12, 8, BLUE)
        shape.click_action.hyperlink.address = SOURCES[key]
    slide.notes_slide.notes_text_frame.text = (
        narrative + "\n\nImplementation inspected 2026-10-21:\n" + IMPLEMENTATION
        + "\n\nCurrent runtime is the FastAPI portal. The separate hosted-agent entrypoint is a "
        "reference scaffold, not the deployed orchestrator. No WorkIQ, M365 delegated identity, "
        "per-user data isolation or video-understanding service is claimed. Browser observations "
        "and model-generated assessments are distinct. Publisher access and terms still apply.\n\n"
        + "Visual reference: " + json.dumps(REFERENCE)
        + "\n\nAzure service icons: Microsoft Azure Public Service Icons V24. Original SVGs and "
        "usage terms are retained in docs/assets/azure-icons. PNGs are proportional transparent "
        "renders; no cropping, recoloring, flipping or rotation. The hosting icon represents the "
        "Container Apps environment, not a Foundry-hosted agent. The general Toolbox icon "
        "represents tool brokerage, not a distinct product logo. No Azure product icon is assigned "
        "to WebIQ, mixed-inspector groups or abstract workflow stages."
        + "\n\nOfficial authoring/vendor references:\n"
        + "\n".join(f"{name}: {url}" for name, url in SOURCES.items())
    )
    return slide


def architecture(prs):
    s = page(prs, "01 / Logical Azure architecture", "From a product link to an evidence-backed decision",
             "The app orchestrates; the model reasons; tools discover and observe; private storage retains permitted evidence.",
             "Dashed groups show logical responsibilities, not a VNet or guaranteed network isolation. "
             "The FastAPI app invokes Foundry for planning/synthesis and owns governed skill loading. "
             "Discovery selects WebIQ Search with a supplied request key or Toolbox web_search without one. "
             "The selected inspector is reused for listing and review pages. Azure-hosted operation uses "
             "managed identity for Foundry and Blob; worker and WebIQ authentication are separate.",
             ("Toolbox", "Blob identity"))
    zone(s, "USER DEVICE", 36, 128, 132, 258, TEAL)
    zone(s, "AZURE CONTAINER APPS ENVIRONMENT", 188, 128, 236, 258, BLUE,
         icon_key="container-apps-environment")
    zone(s, "SERVICE AND TOOL ADAPTERS", 444, 128, 294, 258, BLUE)
    zone(s, "PUBLIC WEBSITES", 758, 128, 166, 258, PURPLE)
    node(s, "Browser", "Portal inputs\nProduct URL\nBuyer priorities\nLive progress\nSaved research",
         48, 172, 108, 158, TEAL, 10.7, icon_key="browser")
    node(s, "FastAPI + SSE portal", "Azure-hosted deployment\nAgent-planned workflow\n\nGoverned SKILL.md\nBounded tool execution\nEvidence checks + gaps", 204, 160, 204, 204)
    node(s, "Foundry Models", "Plan queries; interpret; synthesize", 460, 160, 262, 64,
         icon_key="foundry-models")
    node(s, "Review discovery", "WebIQ Search OR Toolbox web_search", 460, 238, 262, 64, PURPLE, 10.5)
    node(s, "Selected inspector", "WebIQ Browse / BAT / Local / VM", 460, 316, 262, 64, TEAL, 10.8)
    node(s, "Source pages", "Product listing\nReview sites\nForums + retailer Q&A\nPublic video-page text", 774, 172, 134, 184, PURPLE, 10.5)
    route(s, [(156, 231), (204, 231)], TEAL)
    number(s, 1, 180, 216, TEAL)
    for value, y, color in [(2, 192, BLUE), (3, 270, PURPLE), (4, 348, TEAL)]:
        route(s, [(408, y), (460, y)], color)
        number(s, value, 434, y - 15, color)
    route(s, [(722, 270), (774, 270)], PURPLE)
    route(s, [(722, 348), (774, 348)], TEAL)
    route(s, [(204, 298), (156, 298)], TEAL, dashed=True)
    number(s, 6, 180, 283, TEAL)
    node(s, "Azure Storage / Blob", "Generated brief + observations + PNGs", 188, 401, 236, 60, TEAL, 9.8,
         icon_key="storage-account")
    route(s, [(306, 364), (306, 401)], TEAL)
    number(s, 5, 329, 386, TEAL)
    text(s, "2  Model calls     3  Search candidates     4  Observe eligible pages", 460, 400, 464, 20, 10.2, INK)
    text(s, "5  Retain permitted evidence     6  Stream results + reopen via app", 460, 424, 464, 20, 10.2, INK)
    band(s, "Managed identity -> Foundry + private Blob. WebIQ/worker access is separate; website permission is never inherited.",
         icon_key="managed-identity")


def research_flow(prs):
    s = page(prs, "02 / Agent-planned research flow", "Four stable stages; the investigation adapts",
             "Buyer intent shapes the questions and follow-up actions. Search discovery is not the same as reading a page.",
             "The four-stage contract is listing, discovery, optional review inspection and brief. "
             "A search-only run still inspects the initial listing. Extra community exploration is "
             "available only on browser inspection routes (BAT, Local, VM) and within the configured "
             "domain/action/time limits. Review comments are assertions, not authenticated ownership. "
             "The output remains advisory; cart approval is a simulation and checkout is blocked.")
    node(s, "Product URL + priorities", "The user's starting point", 36, 133, 204, 66, BLUE)
    route(s, [(138, 199), (138, 232)])
    items = [
        (36, "1  Listing", "Observe identity, variant,\nprice and seller claims.", BLUE),
        (264, "2  Discovery", "Plan queries; find relevant\nreview/community sources.", PURPLE),
        (492, "3  Inspection", "Read eligible pages;\nrecord text, status and\navailable true PNGs.", TEAL),
        (720, "4  Brief", "Reconcile claims;\nexplain fit, risks and gaps;\nattribute the assessment.", BLUE),
    ]
    for x, title, body, color in items:
        node(s, title, body, x, 232, 204, 122, color, 12)
    for x in (240, 468, 696):
        route(s, [(x, 293), (x + 24, 293)])
    route(s, [(480, 293), (480, 190), (822, 190), (822, 232)], AMBER)
    text(s, "Search-only reviews: skip stage 3", 506, 159, 310, 23, 12, AMBER, True)
    node(s, "Restrictions remain gaps", "Denied or unread does not mean inspected.", 36, 397, 354, 64, AMBER, 11)
    node(s, "Optional bounded exploration", "Allowed domains: search, follow links, scroll.", 420, 397, 318, 64, TEAL, 10.3)
    route(s, [(582, 354), (582, 397)], TEAL)
    route(s, [(612, 397), (612, 354)], TEAL, dashed=True)
    text(s, "Observe / decide / act", 427, 368, 145, 20, 10.5, TEAL)
    node(s, "Human decision", "Cart simulation only;\nnever checkout.", 762, 385, 162, 76, BLUE, 10.5)
    route(s, [(822, 354), (822, 385)])
    band(s, "The model chooses questions and bounded actions; application code enforces access limits, provenance and transaction blocks.")


def request_sequence(prs):
    s = page(prs, "03 / Numbered request and response flow", "One research request, several accountable components",
             "This sequence explains responsibility, not a claim that the whole workflow is one model call.",
             "SSE progress streams throughout the research. Independent evidence is checkpointed "
             "during the run; synthesis adds the final generated brief. Follow-up reads retained "
             "evidence and does not rerun discovery or the model. The portal serves saved objects "
             "through its gate; it does not expose anonymous Blob URLs.", ("Toolbox", "Blob identity"))
    xs = [112, 296, 480, 664, 848]
    actors = [("User", NAVY, None), ("Container Apps", NAVY, "container-apps-environment"),
              ("Foundry Models", BLUE, "foundry-models"), ("Search / inspector", PURPLE, None),
              ("Storage / Blob", TEAL, "storage-account")]
    for x, (label, color, icon_key) in zip(xs, actors):
        rect(s, x - 76, 133, 152, 34, color)
        if icon_key:
            rect(s, x - 71, 137, 26, 26, WHITE)
            service_icon(s, icon_key, x - 68, 140, 20)
        left = x - 39 if icon_key else x - 70
        text(s, label, left, 143, x + 70 - left, 21, 11, WHITE, True, align=PP_ALIGN.CENTER)
        route(s, [(x, 169), (x, 454)], BORDER, dashed=True, arrow=False)
    messages = [
        (0, 1, "URL + intent + method", BLUE, False),
        (1, 3, "Inspect initial product listing", PURPLE, False),
        (3, 1, "Observed identity, claims and available capture", PURPLE, True),
        (1, 2, "Plan buyer-specific queries", BLUE, False),
        (1, 3, "Search; inspect eligible reviews; bounded follow-up", PURPLE, False),
        (3, 1, "Attributed context + access status + eligible PNGs", PURPLE, True),
        (1, 2, "Reconcile claims; state gaps", BLUE, False),
        (1, 4, "Checkpoints during run; final generated brief", TEAL, False),
        (4, 1, "Saved run + retained original evidence", TEAL, True),
        (1, 0, "Brief + protected follow-up", BLUE, True),
    ]
    for index, (src, dst, label, color, dashed) in enumerate(messages, 1):
        y = 190 + (index - 1) * 28
        left, right = sorted((xs[src], xs[dst]))
        text(s, label, left + 5, y - 22, right - left - 10, 22, 10, INK, align=PP_ALIGN.CENTER)
        route(s, [(xs[src], y), (xs[dst], y)], color, dashed=dashed)
        number(s, index, xs[src], y, color)
    band(s, "SSE progress streams throughout steps 2-9. Reopening saved research is a read, not another paid research workflow.")


def routing_choices(prs):
    s = page(prs, "04 / Two independent routing choices", "Five research paths, not five different agents",
             "Choose discovery separately from where and how the listing/review pages are inspected.",
             "Supplying a WebIQ key selects WebIQ Search; an invalid supplied key does not silently "
             "switch to another provider. Without a key, discovery uses the configured Toolbox "
             "web_search. WebIQ Browse itself requires a key. The selected inspector is reused "
             "for listing and reviews when inspection is enabled. Search-only is the separate "
             "inspect_review_pages=false strategy, not an execution_mode. Local and VM require "
             "the cloud model; the native VM worker does not require the desktop extension.",
             ("Toolbox",))
    zone(s, "A  REVIEW DISCOVERY", 36, 134, 414, 197, PURPLE)
    node(s, "WebIQ key?", "Request-scoped choice", 52, 205, 158, 76, PURPLE, 10)
    node(s, "WebIQ Search", "Key supplied", 266, 166, 168, 64, PURPLE, 10.5)
    node(s, "Toolbox web_search", "No key supplied", 266, 252, 168, 64, PURPLE, 10.5,
         icon_key="toolbox")
    route(s, [(210, 228), (238, 228), (238, 198), (266, 198)], PURPLE)
    route(s, [(210, 254), (238, 254), (238, 284), (266, 284)], PURPLE)
    zone(s, "B  CHOOSE ONE INSPECTOR", 474, 134, 450, 197, TEAL)
    for title, body, x, y, icon_key in [
        ("WebIQ Browse", "API text; no viewport PNG", 490, 166, None),
        ("Azure BAT", "Foundry Toolbox browser", 710, 166, "toolbox"),
        ("Local browser", "Visitor device + extension", 490, 252, "browser"),
        ("Azure VM + Playwright", "Windows worker; polls portal", 710, 252, "virtual-machine"),
    ]:
        node(s, title, body, x, y, 198, 64, TEAL, 10, icon_key=icon_key)
    node(s, "C  Review strategy", "Inspect reviews: reuse the selected inspector after discovery.\n"
         "Search-only reviews: inspect the listing, search and synthesize; do not visit review pages.",
         36, 355, 888, 100, BLUE, 13)
    band(s, "Local / VM changes browser placement, not model location. Invalid-key errors do not silently switch discovery providers.")


def evidence_boundary(prs):
    s = page(prs, "05 / Evidence lifecycle and access boundaries", "Transient context is not the evidence archive",
             "Retain the attributed work product and permitted independent observations, not the raw search/Browse response.",
             "Raw WebIQ Search/Browse grounding is active-workflow-only and is not archived. "
             "Generated internal attributed work product and independently obtained browser "
             "observations/PNGs are separate artifact classes and remain subject to applicable "
             "provider and publisher terms. Azure-hosted storage is a private Blob container "
             "authorized with the app identity; local development uses an ignored private directory. "
             "All passcode holders share view/delete access; the prototype does not implement "
             "per-user authorization. Deleting an archive cannot remove already-downloaded ZIPs.",
             ("Blob identity",))
    zone(s, "ACTIVE WORKFLOW", 36, 134, 264, 306, PURPLE)
    zone(s, "PERMITTED DURABLE EVIDENCE", 348, 134, 264, 306, TEAL)
    zone(s, "FOLLOW-UP THROUGH THE APP", 660, 134, 264, 306, BLUE)
    node(s, "Transient retrieval", "Raw WebIQ/Search/Browse\nUsed in this workflow only", 52, 169, 232, 96, PURPLE, 11)
    text(s, "Raw grounding is NOT archived.", 52, 278, 232, 23, 10.5, PURPLE, True)
    node(s, "Independent browser data", "Observed text + access status\nOriginal captured viewports", 52, 316, 232, 100, TEAL, 11)
    node(s, "Generated internal work", "Attributed brief + claim checks\nExplicit evidence gaps", 364, 169, 232, 96, BLUE, 11)
    node(s, "Azure Storage / Blob", "Report + observations + PNGs\nOriginal hashes + timestamps", 364, 316, 232, 100, TEAL, 11,
         icon_key="storage-account")
    node(s, "Saved research", "Passcode-gated follow-up\nShared internal-demo library", 676, 169, 232, 96, BLUE, 11)
    node(s, "Download ZIP", "research.json + report.md\nOriginal PNG files when present", 676, 316, 232, 100, BLUE, 11)
    route(s, [(284, 217), (364, 217)], PURPLE)
    text(s, "Synthesize", 298, 190, 58, 18, 8.8, PURPLE, align=PP_ALIGN.CENTER)
    route(s, [(480, 265), (480, 316)], TEAL)
    route(s, [(284, 366), (364, 366)], TEAL)
    route(s, [(596, 366), (640, 366), (640, 217), (676, 217)], TEAL)
    route(s, [(792, 265), (792, 316)], BLUE)
    text(s, "Until explicit deletion. Downloaded copies are outside archive deletion.", 36, 445, 888, 20, 10.7, GRAY)
    band(s, "Storage needs identity authorization AND a permitted network path. A shared passcode is not per-user isolation.",
         icon_key="managed-identity")


def build(output: Path) -> dict:
    if output.exists():
        raise ValueError("Existing presentations are never overwritten; choose a new output path.")
    icons = validate_icon_assets()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    for draw in (architecture, research_flow, request_sequence, routing_choices, evidence_boundary):
        draw(prs)
    prs.core_properties.title = "Product Due-Diligence Assistant - Solution concept"
    prs.core_properties.subject = "Editable architecture, agent workflow, request sequence, routing and evidence lifecycle"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return {
        "output": str(output), "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "slides": len(prs.slides), "diagram_names": list(DIAGRAMS), "visual_reference": REFERENCE,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "provider_calls": 0, "deployment_changes": False, "diagram_structure_native_editable": True,
        "icon_provenance": {
            "source_page": icons["source_page"], "archive_url": icons["archive_url"],
            "archive_sha256": icons["archive_sha256"], "keys": list(ICON_KEYS),
            "embedded_occurrences": sum(shape.name.startswith("azure-icon-")
                                        for slide in prs.slides for shape in slide.shapes),
            "original_svg_bytes_preserved": True,
        },
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--render-dir", type=Path, help="Optional empty folder; requires desktop Windows PowerPoint.")
    args = parser.parse_args()
    sidecar = args.output.with_suffix(".provenance.json")
    if sidecar.exists() or (args.render_dir and args.render_dir.exists() and any(args.render_dir.iterdir())):
        parser.error("Use new output/metadata paths and an empty render directory.")
    result = build(args.output)
    if args.render_dir:
        result["render"] = render_and_audit(args.output, args.render_dir.resolve())
        result["images"] = []
        for index, name in enumerate(DIAGRAMS, 1):
            image = args.render_dir / (name + ".png")
            (args.render_dir / f"slide-{index:02}.png").rename(image)
            result["images"].append({"file": str(image), "sha256": hashlib.sha256(image.read_bytes()).hexdigest()})
    sidecar.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
