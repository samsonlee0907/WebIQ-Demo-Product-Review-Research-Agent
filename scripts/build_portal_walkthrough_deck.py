"""Build a screenshot-led edition without overwriting the user's Educational deck.

Consumes the receipt from capture_portal_walkthrough.py. Does not call providers.
New walkthrough slides precede the retained technical material; old product
framing slides are replaced and the previous benchmark is labelled historical.
"""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_educational_deck import (
    BLUE, BORDER, GRAY, INK, NAVY, PALE, TEAL, WHITE, box, rect, text, render_and_audit,
)

ROOT = Path(__file__).resolve().parents[1]
ANNOUNCEMENT = "https://blogs.windows.com/devices/2026/10/07/pre-order-our-most-powerful-surface-devices-ever/"
PRODUCT = "https://www.microsoft.com/en-us/surface/devices/surface-laptop-ultra"
PLAYWRIGHT = "https://playwright.dev/python/docs/screenshots"
PPTX_DOCS = "https://python-pptx.readthedocs.io/en/latest/user/slides.html"

STEPS = [
    ("02-inputs", "1 / Paste the product link and buyer priorities",
     "Enter a real URL and the decision you need help making; do not use the built-in example button.",
     "The US example asks about coding, local AI, calls and a US$3,000 budget. Exact configuration and current price remain checks, not assumptions."),
    ("03-method", "2 / Choose the inspector before starting",
     "Select Azure VM + Playwright and confirm the VM worker is online.",
     "This run uses one VM inspector for the listing and eligible review pages, plus Foundry Toolbox web_search. It is not an offline AI model or a five-path benchmark."),
    ("03b-discovery", "Discovery is a separate choice",
     "Leave the WebIQ key empty: the portal uses the operator's configured Foundry Toolbox web_search.",
     "A WebIQ key would select WebIQ Search instead. Choosing the VM inspector does not automatically choose WebIQ, and a browser worker is not the language model."),
    ("04-community", "3 / Enable bounded community follow-up",
     "Open community exploration, turn it on, and enter relevant domains before starting.",
     "The allowlist limits extra agent exploration, not general discovery. Comments are unverified assertions. The model may search, follow observed links, scroll or finish."),
    ("05-running", "4 / Start once and follow the four stages",
     "Click Start research, then follow Listing > Discovery > Inspection > Brief.",
     "This is the actual in-flight progress screen from the single run. A started or completed stage does not imply every source was accessible."),
    ("06-brief", "5 / Read the buying brief first",
     "Read the recommendation, potential fit, watch-outs and evidence gaps before opening detailed claims.",
     "The screenshot shows the generated result of this run, not a scripted success example. Newly launched devices can have sparse independent or owner evidence."),
    ("12-saved", "6 / Reopen the saved follow-up link",
     "Open Saved research or the run's follow-up link; the existing portal passcode is still required.",
     "The generated assessment and independent evidence are retained until explicitly deleted. This is a shared internal-demo library, not per-user authorization."),
    ("07-claims", "7 / Check the claim and its attribution",
     "Compare seller statements with the generated evidence assessment, then follow the source attribution.",
     "A seller's maximum specification is not necessarily the affordable configuration. Attributed synthesis is not a measurement, verified ownership or a certification."),
    ("08-timeline", "8 / Review retained inspection coverage",
     "In Saved research, scroll to Browser evidence: each expandable row identifies an observed page and its status.",
     "This is the saved evidence inventory, not the transient live timeline. Reopening it does not start another model call; a discovered URL is not proof it was read."),
    ("09-listing", "9 / Verify the observed listing baseline",
     "Expand the first Browser evidence row to inspect the listing observation and extraction provenance.",
     "Unknown price, variant or stock remains unknown. Browser observations and true viewport captures are distinct from generated stage illustrations."),
    ("10-viewport", "10 / Open a genuine captured viewport",
     "In the saved screenshot gallery, open an image in a new tab to inspect the retained original.",
     "This is the portal-served original image in the browser, not a synthetic page or the live timeline's image-viewer dialog."),
    ("11-inspection", "11 / Inspect coverage and restrictions",
     "Expand a restricted Browser evidence row to inspect the recorded access status instead of assuming it was read.",
     "A discovered URL is not proof it was read. Review comments are not automatically authentic; access denials and unresolved challenges remain gaps."),
    ("13-retained", "12 / Revisit original screenshots later",
     "Scroll to Genuine browser screenshots and inspect the capture time, source URL and SHA-256 digest.",
     "These are retained original viewport bytes. Raw WebIQ search/Browse grounding is not restored or archived, and screenshots are not generated replacements."),
]


def load_receipt(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("schema") != "portal-walkthrough-v1" or not data.get("finished_at"):
        raise ValueError("Use a completed, genuine capture receipt.")
    if data.get("paid_runs_requested") != 1 or data.get("raw_grounding_saved") is not False:
        raise ValueError("The walkthrough must represent exactly one run without persisted raw grounding.")
    for shot in data["screens"]:
        image = path.parent / shot["file"]
        if image.parent.resolve() != path.parent.resolve() or image.suffix != ".png":
            raise ValueError("Screenshot paths must name adjacent PNG files.")
        if shot.get("kind") != "genuine_portal_screenshot" or shot.get("content_modified") is not False:
            raise ValueError("Screenshots must be unmodified captures of the real portal.")
        content = image.read_bytes()
        if not content.startswith(b"\x89PNG\r\n\x1a\n"):
            raise ValueError(f"Not a PNG screenshot: {shot['file']}")
        if hashlib.sha256(content).hexdigest() != shot["sha256"]:
            raise ValueError(f"Screenshot hash mismatch: {shot['file']}")
    return data


def notes(slide, value, receipt, shot=None):
    body = (
        value + "\n\nLive walkthrough provenance\n"
        f"Product: {receipt['product_url']}\nBuyer intent: {receipt['user_intent']}\n"
        f"Run: {receipt['saved_run'].get('run_id')}\nRevision: {receipt['revision']}\n"
        f"Saved follow-up (authenticated): {receipt['saved_research_url']}\n"
        f"Captured: {receipt['started_at']} to {receipt['finished_at']}\n"
        "Exactly one live research run. Screens show genuine application states; labels and instructions "
        "are editable slide elements outside the screenshots. No credential, cookie, raw SSE body "
        "or raw WebIQ/Search grounding is included in the capture receipt.\n"
        f"\nOfficial product announcement: {ANNOUNCEMENT}\nOfficial product page: {PRODUCT}\n"
        "Announcement checked 2026-10-21: published October 7, availability beginning October 16 "
        "in selected markets. Starting MSRP is not an observed purchasable configuration.\n"
        f"\nScreenshot method: {PLAYWRIGHT}\nPresentation library: {PPTX_DOCS}\n"
    )
    if shot:
        body += (f"\nScreenshot file: {shot['file']}\nSHA-256: {shot['sha256']}\n"
                 f"Capture timestamp: {shot['captured_at']}\nPage: {shot['page_url']}\n"
                 f"Captured region: {shot['region']}; {shot['width']} x {shot['height']} pixels.\n"
                 "Images are proportionally fitted, never stretched or recreated. The region may omit "
                 "other page sections; use the saved run for the complete retained evidence.\n")
    note_slide = slide.notes_slide
    if note_slide.notes_text_frame is None:
        # Edited decks can have an empty notes master; restore only the new slide's body.
        placeholder = next(shape for shape in Presentation().notes_master.placeholders
                           if shape.placeholder_format.type == PP_PLACEHOLDER.BODY)
        note_slide.shapes.clone_placeholder(placeholder)
    note_slide.notes_text_frame.text = body


def new_slide(prs, label, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 960, 6, NAVY, kind=MSO_SHAPE.RECTANGLE)
    text(slide, label.upper(), 36, 22, 888, 19, 10.5, TEAL, True)
    text(slide, title, 36, 49, 888, 72, 27, NAVY, True)
    if subtitle:
        text(slide, subtitle, 36, 121, 888, 42, 13, GRAY)
    return slide


def image_fit(slide, image: Path, x, y, width, height):
    with Image.open(image) as source:
        w, h = source.size
    scale = min(width / w, height / h)
    w, h = w * scale, h * scale
    picture = slide.shapes.add_picture(str(image), Pt(x + (width - w) / 2), Pt(y + (height - h) / 2),
                                      width=Pt(w), height=Pt(h))
    picture.name = "genuine-portal-screenshot"
    return picture


def footer(slide, index, label="LIVE PORTAL WALKTHROUGH | 21 OCT 2026"):
    text(slide, label, 36, 513, 790, 17, 8.5, GRAY)
    text(slide, f"{index:02}", 880, 513, 44, 17, 9, GRAY, align=PP_ALIGN.RIGHT)


def build(base: Path, receipt_path: Path, output: Path) -> dict:
    provenance = output.with_suffix(".provenance.json")
    if (output.resolve() in {base.resolve(), receipt_path.resolve()} or output.exists()
            or provenance.resolve() in {base.resolve(), receipt_path.resolve()} or provenance.exists()):
        raise ValueError("Choose a new output file; existing presentations are never overwritten.")
    receipt = load_receipt(receipt_path)
    screens = {shot["key"]: shot for shot in receipt["screens"]}
    required = {"01-home", "14-export", *(key for key, *_ in STEPS)}
    if not required <= screens.keys():
        raise ValueError("Missing required screenshots: " + ", ".join(sorted(required - screens.keys())))
    base_hash = hashlib.sha256(base.read_bytes()).hexdigest()
    prs = Presentation(base)
    original = list(prs.slides)
    expected = {2: "BUYER-FIRST FRAMING", 6: "SONY WALKTHROUGH / LISTING",
                7: "SONY WALKTHROUGH / QUERY PLANNING"}
    if len(original) < 7 or any(
        label not in "\n".join(shape.text for shape in original[index - 1].shapes if shape.has_text_frame)
        for index, label in expected.items()
    ):
        raise ValueError("The source deck's framing slides changed; review their positions before replacing anything.")
    original_ids = list(prs.slides._sldIdLst)
    for index, slide in enumerate(original, 1):
        slide.part._partname = PackURI(f"/ppt/slides/slide{1000 + index}.xml")
    for sid in original_ids:
        prs.slides._sldIdLst.remove(sid)
    replaced = {2, 6, 7}
    for number in replaced:
        prs.part.drop_rel(original_ids[number - 1].rId)
    inherited = []

    s = new_slide(prs, "A real end-to-end example", "Surface Laptop Ultra\nFrom product link to buying brief")
    text(s, "One live run. Numbered actions. Genuine portal screenshots.", 36, 143, 888, 28, 17, GRAY)
    image_fit(s, receipt_path.parent / screens["01-home"]["file"], 36, 193, 590, 282)
    box(s, f"Follow the first {len(STEPS) + 4} slides", "Paste + prioritize\nChoose + explore\nRead + verify\nSave + export",
        654, 193, 270, 235, TEAL, 18)
    text(s, "Open this saved research (sign-in required)", 654, 438, 270, 35, 12, BLUE,
         True).click_action.hyperlink.address = receipt["saved_research_url"]
    notes(s, "Walkthrough edition based on the user's edited Educational presentation. "
          "The original deck is not overwritten; the technical diagrams are retained in the optional reference chapter.", receipt, screens["01-home"])
    footer(s, len(prs.slides))

    s = new_slide(prs, "Scope and freshness", "Why this laptop for the walkthrough",
                  "Microsoft's own announcement is the freshness check; the portal supplies the buyer-specific investigation.")
    for x, heading, body in [
        (36, "07 OCT 2026", "Official pre-order announcement\nSurface Laptop Ultra"),
        (339, "16 OCT 2026", "Announced availability begins\nSelected markets"),
        (642, "21 OCT 2026", "This captured walkthrough\nUS product URL; US$ budget"),
    ]:
        box(s, heading, body, x, 191, 282, 140, TEAL, 15)
    text(s, "The question", 36, 358, 220, 28, 19, NAVY, True)
    text(s, "Does the evidence support coding, local AI and everyday calls within US$3,000?", 36, 394, 888, 58, 22, INK, True)
    text(s, "Microsoft announcement / availability dates", 36, 463, 420, 24, 11, BLUE,
         True).click_action.hyperlink.address = ANNOUNCEMENT
    text(s, "Open the official product page", 590, 463, 334, 24, 11, BLUE,
         True).click_action.hyperlink.address = PRODUCT
    text(s, "Maximum specifications and starting MSRP are not a verified configuration.", 36, 489, 888, 19, 10.5, GRAY)
    notes(s, "The directly retrieved Microsoft Devices blog and official product page establish the product "
          "and published availability date. This is not a claim that every region or configuration is in stock. "
          "The example does not submit an order or operate checkout.", receipt)
    footer(s, len(prs.slides))

    for key, title, instruction, explanation in STEPS:
        shot = screens[key]
        s = new_slide(prs, "Follow along / actual portal screen", title, instruction)
        rect(s, 36, 178, 888, 291, PALE, BORDER)
        image_fit(s, receipt_path.parent / shot["file"], 43, 183, 874, 281)
        text(s, explanation, 36, 477, 888, 30, 10.5, GRAY)
        notes(s, instruction + "\n\n" + explanation, receipt, shot)
        footer(s, len(prs.slides))

    s = new_slide(prs, "13 / Keep a usable follow-up", "Export the report and genuine evidence",
                  "Click Download evidence + screenshots. Keep the protected follow-up link separately.")
    image_fit(s, receipt_path.parent / screens["14-export"]["file"], 36, 178, 888, 77)
    for x, heading, body in [
        (36, "research.json", "Run metadata\nAttributed assessment\nIndependent observations"),
        (339, "report.md", "Readable generated brief\nSource attribution\nExplicit evidence gaps"),
        (642, "screenshots/*.png", "Original captured viewports\nSHA-256 identity\nNo generated replacements"),
    ]:
        box(s, heading, body, x, 285, 282, 155, TEAL, 15)
    verified = receipt["export_verified"]
    text(s, f"Actual export verified: {verified['genuine_pngs']} genuine PNGs; original hashes matched. "
         "Deletion does not remove copies you already downloaded.", 36, 463, 888, 40, 12, GRAY)
    notes(s, "The download action was exercised. Every original screenshot's SHA-256 was compared with "
          "the export manifest. The ZIP was inspected in temporary browser download storage and is not "
          "included in the public source. Shared passcode access is not confidential multi-user isolation.", receipt, screens["14-export"])
    footer(s, len(prs.slides))

    saved = receipt["saved_run"]
    s = new_slide(prs, "What this one run actually established", "Keep the outcome and uncertainty visible",
                  "This is an observed workflow example, not a controlled product test or a new five-path ranking.")
    for x, heading, body in [
        (36, str(saved["genuine_screenshots"]), "Genuine retained viewport PNGs\nNot the number of slide screenshots"),
        (339, str(saved["claim_checks"]), "Generated claim checks\nOpen the attribution"),
        (642, str(saved["community_actions"]), "Recorded exploration decisions\nNot authenticated owner reports"),
    ]:
        box(s, heading, body, x, 184, 282, 124, TEAL, 14)
    text(s, f"Observed outcome: {receipt['ui_outcome']}", 36, 333, 888, 36, 22, NAVY, True)
    summary = saved.get("executive_summary") or "No completed generated assessment was available."
    text(s, summary, 36, 378, 888, 94, 14, INK)
    text(s, "Follow the evidence gaps in the saved report. No purchase or real cart mutation was attempted.", 36, 480, 888, 26, 11.5, GRAY)
    notes(s, "Actual saved-run summary:\n" + json.dumps(saved, indent=2) +
          "\nThe VM observation count and capture count are different measures. A completed synthesis "
          "does not establish price eligibility, exhaustive review coverage or long-term reliability.", receipt)
    footer(s, len(prs.slides))
    walkthrough_count = len(prs.slides)

    s = new_slide(prs, "Optional reference chapter", "How the workflow is implemented",
                  "Retained diagrams from the edited Educational deck explain services, placement, evidence and governance.")
    for x, title, body in [
        (36, "Agent and orchestration", "Planned questions\nGoverned skills\nBounded tool execution"),
        (339, "Discovery and inspection", "WebIQ / Toolbox search\nBrowse / BAT / Local / VM\nExplicit search-only strategy"),
        (642, "Evidence and limits", "Original captures\nAttribution and restrictions\nHuman accountability"),
    ]:
        box(s, title, body, x, 211, 282, 180, TEAL, 15)
    text(s, "The following historical comparison uses a different product. Its numbers are not Surface Laptop Ultra results.",
         36, 447, 888, 48, 15, GRAY)
    notes(s, "Technical material is retained from the base presentation, except the old product-framing "
          "slides 2, 6 and 7 replaced by this real walkthrough. Historical comparison slides remain labelled separately.", receipt)
    footer(s, len(prs.slides))

    for index, (slide, sid) in enumerate(zip(original, original_ids), 1):
        if index in replaced:
            continue
        if index == 29:
            divider = new_slide(prs, "Historical comparison / different product", "Do not confuse the benchmark with this walkthrough",
                                "The following four recorded comparison slides predate the laptop run and concern the earlier example.")
            text(divider, "Different product\nDifferent research round\nNot a laptop recommendation", 36, 220, 888, 165, 31, NAVY, True)
            notes(divider, "The inherited five-path timing/coverage/source metrics belong to the earlier "
                  "recorded example, not the single Microsoft laptop workflow captured here.", receipt)
            footer(divider, len(prs.slides))
        sid.id = max((item.id for item in prs.slides._sldIdLst), default=255) + 1
        prs.slides._sldIdLst.append(sid)
        for shape in slide.shapes:
            if (index == 16 and shape.has_text_frame
                    and shape.text == "Cloud-native Browser Automation Tool" and shape.height < Pt(48)):
                shape.height = Pt(48)
            if shape.has_text_frame and shape.text.strip().isdigit() and shape.left > Pt(850) and shape.top > Pt(500):
                shape.text_frame.paragraphs[0].runs[0].text = f"{len(prs.slides):02}"
        inherited.append({"source_slide": index, "new_slide": len(prs.slides)})
    prs.core_properties.title = "Surface Laptop Ultra - Portal walkthrough and technical reference"
    prs.core_properties.subject = "One genuine product-research run with numbered portal screenshots"
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    if hashlib.sha256(base.read_bytes()).hexdigest() != base_hash:
        raise RuntimeError("The source presentation changed during generation.")
    summary = {
        "output": str(output), "slides": len(prs.slides), "walkthrough_slides": walkthrough_count,
        "source_sha256": base_hash, "source_slides": len(original), "retained_slides": inherited,
        "replaced_product_framing_slides": sorted(replaced),
        "receipt_sha256": hashlib.sha256(receipt_path.read_bytes()).hexdigest(),
        "embedded_screenshot_count": len(screens), "saved_research_url": receipt["saved_research_url"],
        "generated_at": datetime.now().astimezone().isoformat(),
    }
    provenance.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", type=Path, default=ROOT / "docs" / "Product-Review-Agent-Discussion-Educational.pptx")
    parser.add_argument("--receipt", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=ROOT / "docs" / "Product-Review-Agent-Surface-Laptop-Walkthrough.pptx")
    parser.add_argument("--render-dir", type=Path)
    args = parser.parse_args()
    result = build(args.base, args.receipt, args.output)
    if args.render_dir:
        result["render"] = render_and_audit(args.output, args.render_dir)
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
