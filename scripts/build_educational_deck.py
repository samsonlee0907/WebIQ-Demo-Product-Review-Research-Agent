"""Build a diagram-led educational edition from the user's edited discussion deck.

The expanded presentation is read, never saved or opened through PowerPoint.
Six useful user-edited slides are retained with updated notes/footers; table-heavy
material is retaught with editable native diagrams. No publisher grounding or
private captures are embedded. This generator makes no service/model calls.

Final example:
    python scripts\\build_educational_deck.py --run-not-before 2026-10-21T01:00:00Z

Use --research-only to preview instructional slides before a fresh round exists.
Optional --render-dir exports PNGs and a PowerPoint text-bounds audit. Only the
specific read-only educational presentation opened by this script is closed.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import urlsplit

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "docs" / "Product-Review-Agent-Discussion-Expanded.pptx"
OUTPUT = ROOT / "docs" / "Product-Review-Agent-Discussion-Educational.pptx"
RUNS = ROOT / "docs" / "browser-benchmark-runs.json"
VERIFIED = "2026-10-21"
NAVY, BLUE, CYAN, TEAL = "0B1E3F", "0078D4", "50E6FF", "007A6E"
PURPLE, AMBER, INK, GRAY = "5C2D91", "986A00", "201F1E", "605E5C"
WHITE, PALE, BORDER, MUTED = "FFFFFF", "F3F7FC", "DDE4EE", "8A8886"
METHODS = (
    ("webiq_browse", "WebIQ Browse", BLUE),
    ("live", "Azure BAT Tool", PURPLE),
    ("client_browser", "Local browser", TEAL),
    ("vm_playwright", "VM + Playwright", AMBER),
    ("search_only", "Search-only reviews", NAVY),
)
SOURCES = {
    "W1": ("WebIQ: passages and content formats", "https://learn.microsoft.com/en-us/web-iq/concepts/content-formats"),
    "W2": ("WebIQ: Web Search reference", "https://learn.microsoft.com/en-us/web-iq/reference/web-search"),
    "W3": ("WebIQ: Browse reference", "https://learn.microsoft.com/en-us/web-iq/reference/browse"),
    "W4": ("WebIQ: FAQ", "https://learn.microsoft.com/en-us/web-iq/resources/faq"),
    "W5": ("WebIQ: terms and display/storage limits", "https://www.microsoft.com/en-us/webiq/webiq-legal"),
    "F1": ("Foundry: Browser Automation", "https://learn.microsoft.com/en-us/azure/foundry/agents/how-to/tools/browser-automation"),
    "F2": ("Foundry: Web Search", "https://learn.microsoft.com/en-us/azure/foundry/agents/how-to/tools/web-search"),
    "F3": ("Foundry: create and manage a toolbox", "https://learn.microsoft.com/en-us/azure/foundry/agents/how-to/tools/toolbox"),
    "P1": ("Playwright: Python CDP connection", "https://playwright.dev/python/docs/api/class-browsertype#browser-type-connect-over-cdp"),
    "P2": ("Playwright: network", "https://playwright.dev/python/docs/network"),
    "P3": ("Playwright: screenshots", "https://playwright.dev/python/docs/screenshots"),
    "P4": ("Playwright: browser contexts / isolation", "https://playwright.dev/python/docs/browser-contexts"),
    "P5": ("Playwright: official CDP source documentation", "https://raw.githubusercontent.com/microsoft/playwright/main/docs/src/api/class-browsertype.md"),
    "C1": ("Chrome extensions: messaging", "https://developer.chrome.com/docs/extensions/develop/concepts/messaging"),
    "C2": ("Chrome extensions: scripting", "https://developer.chrome.com/docs/extensions/reference/api/scripting"),
    "C3": ("Chrome extensions: captureVisibleTab", "https://developer.chrome.com/docs/extensions/reference/api/tabs#method-captureVisibleTab"),
    "A1": ("Azure: virtual machines overview", "https://learn.microsoft.com/en-us/azure/virtual-machines/overview"),
    "A2": ("Microsoft Entra: managed identities", "https://learn.microsoft.com/en-us/entra/identity/managed-identities-azure-resources/overview"),
    "L1": ("LangGraph: Graph API / state, nodes, edges", "https://docs.langchain.com/oss/python/langgraph/graph-api"),
    "L2": ("LangGraph: official repository", "https://github.com/langchain-ai/langgraph"),
    "K1": ("Agent Skills: official specification", "https://agentskills.io/specification"),
    "APP": ("Prototype implementation, not vendor guarantees", "api\\app\\pipeline.py; api\\app\\bat_research.py; api\\app\\config.py; browser_extension\\; scripts\\windows_playwright_worker.py"),
    "RUN": ("Fresh operational receipts, not publisher content", "docs\\browser-benchmark-runs.json"),
    "EDIT": ("User-edited source presentation", "docs\\Product-Review-Agent-Discussion-Expanded.pptx"),
}


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=PALE, line=None, kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    shape = slide.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.name = "diagram-box"
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    else:
        shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size=14, fill=INK, bold=False, mono=False, align=None):
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    shape.name = "educational-text"
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for index, line in enumerate(str(value).split("\n")):
        p = frame.paragraphs[0] if not index else frame.add_paragraph()
        p.text = line
        p.font.name = "Consolas" if mono else ("Segoe UI Semibold" if bold else "Segoe UI")
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(fill)
        p.space_after = Pt(5)
        p.line_spacing = 1.05
        if align is not None:
            p.alignment = align
    return shape


def link(slide, x1, y1, x2, y2, fill=BLUE, dashed=False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    shape.name = "diagram-arrow"
    shape.line.color.rgb = rgb(fill)
    shape.line.width = Pt(1.8)
    ln = shape._element.spPr.get_or_add_ln()
    end = OxmlElement("a:tailEnd")
    end.set("type", "triangle")
    ln.append(end)
    if dashed:
        dash = OxmlElement("a:prstDash")
        dash.set("val", "dash")
        ln.append(dash)
    return shape


def box(slide, title, body, x, y, w, h=130, accent=BLUE, size=14):
    rect(slide, x, y, w, h, PALE, BORDER)
    rect(slide, x, y, 5, h, accent, kind=MSO_SHAPE.RECTANGLE)
    compact = h <= 110
    title_h = 29 if compact and body else (h - 23 if not body else 40)
    text(slide, title, x + 15, y + (9 if compact else 12), w - 28, title_h, 16, accent, True)
    if body:
        text(slide, body, x + 15, y + (39 if compact else 59), w - 28,
             h - (45 if compact else 68), size)


def band(slide, value):
    rect(slide, 36, 450, 888, 36, NAVY, kind=MSO_SHAPE.RECTANGLE).name = "layout-band"
    text(slide, value, 48, 459, 864, 21, 10.7, WHITE, True)


def note(slide, narrative, refs=("APP",)):
    if slide.notes_slide.notes_text_frame is None:
        slide.notes_slide.shapes._spTree.append(parse_xml(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<p:nvSpPr><p:cNvPr id="2" name="Educational source notes"/>'
            '<p:cNvSpPr/><p:nvPr><p:ph type="body" idx="3"/></p:nvPr></p:nvSpPr>'
            '<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
        ))
    slide.notes_slide.notes_text_frame.text = (
        narrative
        + "\n\nCapability labels: vendor documentation describes service capabilities; APP describes "
        "implemented behavior; future ideas are explicitly unshipped. Synthetic examples are authored "
        "for teaching, not actual publisher prose or measured Sony product findings.\n\n"
        + "\n\n".join(
            f"[{key}] {SOURCES[key][0]}\n{SOURCES[key][1]}\n"
            + (f"Verified by direct official-page retrieval {VERIFIED}."
               if key not in {"APP", "RUN", "EDIT"} else f"Local source read {VERIFIED}.")
            for key in refs
        )
    )


def page(prs, section, title, subtitle="", refs=("APP",)):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 960, 540, WHITE, kind=MSO_SHAPE.RECTANGLE).name = "layout-background"
    rect(s, 0, 0, 9, 540, BLUE, kind=MSO_SHAPE.RECTANGLE).name = "layout-accent"
    text(s, section.upper(), 36, 20, 888, 19, 10.5, BLUE, True)
    text(s, title, 36, 44, 888, 39, 25, NAVY, True)
    if subtitle:
        text(s, subtitle, 36, 89, 888, 35, 13, GRAY)
    footer = "; ".join(f"{k} {SOURCES[k][0].split(':')[0]}" for k in refs)
    text(s, footer, 36, 493, 888, 16, 8.1, MUTED)
    text(s, f"Educational edition | Official docs checked {VERIFIED}", 36, 515, 800, 14, 8, MUTED)
    text(s, f"{len(prs.slides):02}", 880, 513, 44, 17, 9, MUTED, align=PP_ALIGN.RIGHT)
    note(s, subtitle, refs)
    return s


def row_flow(s, items, y=165, h=165, x=36, width=888):
    gap = 24
    w = (width - gap * (len(items) - 1)) / len(items)
    for i, (title, body, accent) in enumerate(items):
        xx = x + i * (w + gap)
        box(s, title, body, xx, y, w, h, accent, 13.5)
        if i:
            link(s, xx - gap + 3, y + h / 2, xx - 4, y + h / 2)


def sequence(s, headers, messages):
    """Messages are (sender-index, receiver-index, label, color); top-to-bottom time."""
    xs = [100 + i * 760 / (len(headers) - 1) for i in range(len(headers))]
    for x, title in zip(xs, headers):
        rect(s, x - 64, 133, 128, 47, NAVY)
        text(s, title, x - 57, 143, 114, 32, 12.5, WHITE, True, align=PP_ALIGN.CENTER)
        link(s, x, 184, x, 429, BORDER, True)
    for i, (a, b, label, accent) in enumerate(messages):
        y = 215 + i * 47
        left, right = sorted((xs[a], xs[b]))
        link(s, xs[a], y + 12, xs[b], y + 12, accent)
        rect(s, (left + right) / 2 - 154, y - 9, 308, 22, WHITE, kind=MSO_SHAPE.RECTANGLE)
        text(s, label, (left + right) / 2 - 150, y - 8, 300, 21, 11.3, accent, True,
             align=PP_ALIGN.CENTER)


def iso_time(value):
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def fresh_runs(path, not_before):
    """Reject the mixed historical report rather than silently re-labeling it fresh."""
    data = json.loads(path.read_text(encoding="utf-8"))
    selected = {}
    for method, _, _ in METHODS:
        candidates = [
            r for r in data["runs"]
            if r.get("method") == method and r.get("case_id") == "sony-wh1000xm6-official"
            and iso_time(r["started_at"]) >= not_before
        ]
        if not candidates:
            raise ValueError(f"No fresh {method} receipt at/after {not_before.isoformat()}. "
                             "Wait for the finalized five-path round; do not use historical runs.")
        selected[method] = max(candidates, key=lambda r: iso_time(r["started_at"]))
    for field in ("product_url", "intent", "web_query", "server_revision", "source_commit", "synthesis_model"):
        vals = {selected[m].get(field) for m, _, _ in METHODS}
        if len(vals) != 1 or None in vals:
            raise ValueError(f"Fresh receipts must have matching nonempty {field}: {vals}")
    if selected["search_only"].get("inspect_review_pages") is not False:
        raise ValueError("Search-only receipt must explicitly disable review-page inspection.")
    for r in selected.values():
        for key in ("elapsed_s", "discovered", "inspected", "real_captures", "run_id", "completed_at"):
            if key not in r:
                raise ValueError(f"{r['method']} lacks operational metric {key}")
    return data, selected


class Edition:
    def __init__(self, base):
        self.base_hash = hashlib.sha256(base.read_bytes()).hexdigest()
        self.prs = Presentation(base)
        if abs(self.prs.slide_width.pt - 960) > 0.1 or abs(self.prs.slide_height.pt - 540) > 0.1:
            raise ValueError("The edited deck must retain its 960 x 540 point canvas.")
        keep = {2, 16, 24, 26, 27, 30}
        self.retained = {}
        ids = self.prs.slides._sldIdLst
        for n, sid in enumerate(list(ids), 1):
            if n in keep:
                self.retained[n] = (sid, self.prs.slides[n - 1])
                self.prs.slides[n - 1].part._partname = PackURI(f"/ppt/slides/slide{900 + n}.xml")
        for n, sid in enumerate(list(ids), 1):
            ids.remove(sid)
            if n not in keep:
                self.prs.part.drop_rel(sid.rId)
        self.inherited = []

    def retain(self, number, refs, narrative):
        sid, slide = self.retained[number]
        sid.id = max((item.id for item in self.prs.slides._sldIdLst), default=255) + 1
        self.prs.slides._sldIdLst.append(sid)
        self.inherited.append({"source_slide": number, "new_slide": len(self.prs.slides)})
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            if sh.text.startswith("Sources:"):
                sh.text_frame.paragraphs[0].runs[0].text = "Sources: " + ", ".join(refs)
            elif "OCT 2026 | Product-review research" in sh.text:
                sh.text_frame.paragraphs[0].runs[0].text = "21 OCT 2026 | Educational edition"
            elif sh.left.pt > 870 and sh.top.pt > 505:
                sh.text_frame.paragraphs[0].runs[0].text = f"{len(self.prs.slides):02}"
            elif sh.top.pt < 45 and re.match(r"^\d+\s+\|", sh.text):
                sh.text_frame.paragraphs[0].runs[0].text = re.sub(r"^\d+\s+\|\s*", "", sh.text)
        note(slide, narrative + f"\nRetained editable shapes from user-edited source slide {number}.", refs)
        return slide


def build(output=OUTPUT, runs_path=RUNS, not_before=None, research_only=False):
    not_before = not_before or iso_time("2026-10-21T01:00:00Z")
    data, runs = (None, None) if research_only else fresh_runs(runs_path, not_before)
    edition = Edition(BASE)
    p = edition.prs
    props = p.core_properties
    props.title = "Product reviews, understood by agents — educational edition"
    props.subject = "Tools, evidence boundaries and five-path operational walkthrough"
    props.author = "Product Due-Diligence Prototype"
    props.created = props.modified = datetime.now(timezone.utc)
    props.keywords = "educational; editable diagrams; WebIQ; Foundry; Playwright; five paths"
    props.comments = f"Based on the user-edited expanded deck. Source SHA256: {edition.base_hash}"

    s = page(p, "Architecture & engineering tutorial", "Product reviews, understood by agents",
             "One product URL. A buyer's real priorities. Five paths through one explainable evidence journey.",
             ("EDIT", "APP"))
    row_flow(s, [("01  Understand", "Observe the listing.\nSeparate facts from claims.", BLUE),
                 ("02  Discover", "Plan buyer-specific queries.\nFind public sources.", TEAL),
                 ("03  Inspect", "Visit review pages — or\nexplicitly choose not to.", PURPLE),
                 ("04  Explain", "Cite the evidence.\nKeep uncertainty visible.", NAVY)], y=155, h=181)
    text(s, "Sony WH-1000XM6  |  A guided case, not a product endorsement", 40, 371, 880, 28, 18, NAVY, True)
    band(s, "DRAFT RESEARCH PREVIEW — no fresh measured results included." if research_only
         else "Tools retrieve and observe. Models interpret. Code enforces boundaries. The buyer still decides.")

    edition.retain(2, ("APP", "EDIT"), "Preserved buyer-first framing and the full Sony intent. The supplied URL "
                   "does not prove current price, stock or variant. No new product facts are asserted.")

    s = page(p, "The route", "One scaffold; two independent choices",
             "Choose discovery separately from inspection. Search-only changes the review-visit step.")
    row_flow(s, [("Listing", "Seed URL → observed identity,\noffer and seller claims", BLUE),
                 ("Discovery", "Buyer priorities → planned\nqueries → cited candidates", TEAL),
                 ("Review evidence", "Same inspector as listing,\nOR explicit search-only skip", PURPLE),
                 ("Synthesis", "Buyer-fit brief, citations,\nclaim checks and gaps", NAVY)], y=163, h=168)
    rect(s, 45, 357, 552, 58, PALE, BORDER)
    text(s, "Inspector: Browse / BAT / local / VM\nSearch-only defaults to Browse for the listing only.", 59, 367, 528, 43, 13)
    rect(s, 617, 357, 307, 58, PALE, BORDER)
    text(s, "Discovery: WebIQ Search with key\nor Foundry Toolbox web_search", 631, 367, 280, 43, 13)
    band(s, "No hidden inspector switch. Cart simulation is a separate, non-purchasing branch.")
    note(s, "Current app: both listing and optional review visits use the selected inspector. Search-only "
         "reviews still inspect the listing; the user may explicitly choose a non-WebIQ seed inspector "
         "to use Toolbox search without a WebIQ key. Query planning is model-driven; routing is code.", ("APP",))

    s = page(p, "Vocabulary map", "Each component answers a different question",
             "A search service, a model, a browser protocol, a host and a skill are not synonyms.",
             ("W4", "F3", "P5", "K1", "APP"))
    box(s, "MODEL", "What should we ask?\nHow does evidence affect this buyer?", 36, 145, 274, 123, PURPLE)
    box(s, "SEARCH", "Which public sources might help?\nReturns URLs + grounded material.", 343, 145, 274, 123, TEAL)
    box(s, "INSPECTOR", "What can this URL show us?\nReturns text / status / permitted PNG.", 650, 145, 274, 123, BLUE)
    box(s, "ORCHESTRATION", "Which step runs next?\nEnforces state, limits and safety.", 36, 292, 274, 123, NAVY)
    box(s, "HOST + PROTOCOL", "Where is the browser?\nHow do commands reach it?", 343, 292, 274, 123, AMBER)
    box(s, "SKILL", "What instructions apply now?\nGuidance is not an executed tool.", 650, 292, 274, 123, PURPLE)
    band(s, "Keep these categories separate when explaining the architecture, permissions or measured results.")

    s = page(p, "Model vs code", "The model reasons; the scaffold enforces",
             "GPT-5.6-Terra is the configured deployment name — not a claim about public model capabilities.")
    for y, label, accent in [(149, "MODEL", PURPLE), (298, "CODE", BLUE)]:
        rect(s, 36, y, 110, 117, accent)
        text(s, label, 49, y + 43, 84, 29, 16, WHITE, True)
    box(s, "Plan", "Turn the buyer's needs into\nfocused search questions.", 174, 149, 330, 117, PURPLE)
    box(s, "Interpret", "Reconcile cited evidence.\nExplain fit, conflicts and gaps.", 540, 149, 384, 117, PURPLE)
    box(s, "Execute", "Normalize URLs; use the chosen\ninspector; respect deadlines.", 174, 298, 330, 117, BLUE)
    box(s, "Constrain", "Keep provenance, delete grounding,\nblock checkout and label skips.", 540, 298, 384, 117, BLUE)
    link(s, 338, 270, 338, 293)
    link(s, 729, 293, 729, 270)
    band(s, "A well-formed model response is still not independent factual verification.")

    s = page(p, "Orchestration explained", "LangGraph: state + nodes + edges",
             "The listing subflow uses this graph pattern; the broader four-stage route also uses explicit app code.",
             ("L1", "L2", "APP"))
    row_flow(s, [("Setup", "Select the driver.\nCreate/connect session.", BLUE),
                 ("Observe", "Read current page state.\nRecord access and fields.", TEAL),
                 ("Decide", "Enough evidence?\nAnother bounded step?", PURPLE),
                 ("Finalize", "Normalize observation.\nRelease session resources.", NAVY)], y=155, h=145)
    box(s, "Act", "Allowed, bounded browser step", 347, 348, 266, 76, AMBER, 12)
    link(s, 592, 302, 592, 344, PURPLE)
    text(s, "continue", 611, 318, 130, 19, 11.5, PURPLE)
    link(s, 348, 383, 286, 383, AMBER)
    link(s, 286, 383, 286, 306, AMBER)
    text(s, "State carries observed fields, status and remaining budget — not a model's hidden reasoning.",
         36, 427, 888, 19, 11.5, GRAY)
    band(s, "Nodes do work; edges choose the next node; state carries only the data the workflow needs.")
    note(s, "LangGraph defines state, nodes and fixed/conditional edges. Local listing graph contains setup, "
         "observe, decide, act and finalize; do not imply the entire app is one durable LangGraph deployment. "
         "Vendor persistence/human-loop features are not automatically enabled by using StateGraph.",
         ("L1", "L2", "APP"))

    s = page(p, "Skills explained", "Progressive disclosure is about instructions",
             "A SKILL.md packages guidance; an executable tool performs an operation.",
             ("K1", "APP"))
    row_flow(s, [("1  Metadata", "Name + description\nhelp decide relevance.", BLUE),
                 ("2  SKILL.md", "Load the applicable\nstage's instructions.", PURPLE),
                 ("3  Resources", "Read supporting references\nor run a script only if needed.", TEAL)], y=151, h=157)
    box(s, "In this app", "Listing digestion → review discovery → claim-reality reasoning → human-gated cart simulation",
        36, 337, 554, 88, NAVY, 12.5)
    box(s, "Not magic execution", "Loading a skill does not grant credentials,\ncreate a browser or verify a claim.", 617, 337, 307, 88, AMBER, 12)
    band(s, "The app explicitly loads stage guidance. A separate runtime must implement and authorize the tools.")

    s = page(p, "Sony walkthrough / listing", "Build a baseline, not a guess",
             "A supplied model name is input. Successful observation is evidence.")
    box(s, "PUBLIC PRODUCT URL", "electronics.sony.com\nSony WH-1000XM6 seed", 36, 194, 250, 145, BLUE)
    link(s, 291, 266, 337, 266)
    for y, title, body, accent in [
        (139, "Identity", "Observed model / SKU / variant", BLUE),
        (234, "Offer", "Observed seller / currency / price", TEAL),
        (329, "Claims + access", "Seller statements / status / provenance", PURPLE),
    ]:
        box(s, title, body, 350, y, 335, 83, accent, 12.5)
    box(s, "UNKNOWN", "If the page is blocked:\nkeep the missing baseline.\n\nDo not infer a qualifying\nprice under $450.", 717, 169, 207, 211, AMBER, 13)
    band(s, "The baseline anchors queries and claim checks. It never turns seller claims into verified product performance.")

    s = page(p, "Sony walkthrough / query planning", "Translate needs into answerable questions",
             "Synthetic teaching examples — these are research questions, not Sony findings.")
    for i, (need, question, source, accent) in enumerate([
        ("Glasses + 10-hour flights", "Does seal or comfort change\nwith thin frames over time?", "Lab conditions + owner context", BLUE),
        ("Windows + iPhone", "How does multipoint hand-off\nwork for this exact setup?", "Device / OS / firmware tests", TEAL),
        ("Folding hinge + budget", "What is known about durability?\nWas an eligible offer observed?", "Long-term reports + current listing", PURPLE),
    ]):
        y = 143 + i * 96
        box(s, need, "", 36, y, 242, 80, accent)
        link(s, 283, y + 40, 314, y + 40, accent)
        text(s, question, 330, y + 9, 332, 59, 15)
        text(s, source, 690, y + 13, 234, 54, 13, accent, True)
    band(s, "Relevance selects the platform. Reddit, Amazon and YouTube are possible sources, not mandatory templates.")

    s = page(p, "Discovery choice", "Two discovery interfaces; no silent key fallback",
             "The request's credentials determine the search route, independently of the inspection environment.",
             ("W2", "F2", "F3", "APP"))
    box(s, "Planned query", "Buyer priorities + observed identity", 36, 234, 250, 116, NAVY)
    link(s, 290, 264, 363, 194)
    link(s, 290, 314, 363, 366)
    box(s, "WebIQ Search", "Request-scoped key supplied.\nPublic search + attributed passages.", 374, 139, 359, 125, BLUE)
    box(s, "Toolbox web_search", "No WebIQ key supplied.\nFoundry-connected discovery.", 374, 302, 359, 125, TEAL)
    box(s, "Key fails?", "Show the failure.\nDo not quietly switch\nto another provider.", 760, 164, 164, 214, AMBER, 13)
    band(s, "Search-only + default Browse seed still needs WebIQ access; explicitly choose another seed inspector if needed.")

    s = page(p, "What is WebIQ Search?", "A query-to-evidence service, not an LLM",
             "Retrieval and ranking execute in Microsoft infrastructure; this app calls the public HTTPS API.",
             ("W1", "W2", "W4", "APP"))
    row_flow(s, [("Input", 'query + maxResults\ncontentFormat="passage"\nmaxLength=3000', BLUE),
                 ("WebIQ service", "Find relevant public pages.\nSelect query-targeted parts\nof structured documents.", TEAL),
                 ("Output", "URL + title + content\nOptional crawl/update dates\nAttribution for synthesis", PURPLE)], y=150, h=181)
    text(s, "Strength: relevant context with less off-topic material.", 40, 357, 870, 28, 17, TEAL, True)
    text(s, "Limit: a passage is not a new browser visit, a complete document or an independently verified fact.",
         40, 393, 870, 37, 14)
    band(s, "The app requests up to 3,000 characters per source for model context, not merely the 200-character UI excerpt.")

    s = page(p, "Passages / display / retention", "Three limits, three different purposes",
             "Do not confuse model grounding, visible quotation and post-run evidence.",
             ("W1", "W2", "W5", "APP"))
    row_flow(s, [("Up to 3,000", "CHARACTERS / SOURCE\nRequested search passage;\nmodel receives bounded context.", BLUE),
                 ("At most 200", "RAW CHARACTERS / SOURCE\nVisible excerpt ceiling, or lower\nif required by the publisher.", PURPLE),
                 ("Delete on finish", "RAW WEBIQ GROUNDING\nTransient workflow input;\nnot a replayable source archive.", AMBER)], y=151, h=183)
    rect(s, 36, 362, 888, 61, PALE, BORDER)
    text(s, "Expanded text / JSON panels are safe projections and excerpts.\nAn “expand” control is not permission to dump raw API payloads.", 51, 373, 858, 45, 14)
    band(s, "Attribution stays attached to generated output under the applicable terms. This is an implementation boundary, not legal advice.")
    note(s, "Official content-format documentation recommends passage and maxLength 3000. Without full-document "
         "access Search caps content at 3000 and text/html/markdown requests may return 403. App uses passage-3000 "
         "internally, raw excerpts <=200 for display. Terms require raw grounding deletion after the Atomic Action; "
         "summary/associated attribution exceptions are distinct. No publisher content is embedded here.",
         ("W1", "W2", "W5", "APP"))

    s = page(p, "What is Foundry Toolbox search?", "A managed tool interface for web discovery",
             "Toolbox packages reusable tools behind a managed MCP endpoint; it is not itself a browser.",
             ("F2", "F3", "APP"))
    sequence(s, ["App", "Foundry Toolbox", "web_search"], [
        (0, 1, "Authenticate + call discovered tool", BLUE),
        (1, 2, "Submit the public-web question", TEAL),
        (2, 1, "Grounded result + source references", TEAL),
        (1, 0, "Normalize candidates for this workflow", BLUE),
    ])
    band(s, "No WebIQ request key is needed for this route. Foundry identity, access, billing and terms still apply.")
    note(s, "Current Foundry docs distinguish direct WebSearchTool and toolbox WebSearchToolboxTool interfaces. "
         "The configured app uses Toolbox web_search as its alternative discovery provider, not the WebIQ REST "
         "endpoint or its passage/maxLength contract. Search uses Bing grounding under its own applicable terms "
         "and data-handling boundaries. The app's explicit inspector still determines URL observation.",
         ("F2", "F3", "APP"))

    s = page(p, "What is WebIQ Browse?", "A selected URL becomes structured document text",
             "Server-side extraction, not ownership of an interactive browser or a viewport screenshot.",
             ("W3", "W4", "APP"))
    row_flow(s, [("Known URL", "The app supplies the exact\nlisting or review URL.", BLUE),
                 ("Index / crawl", "none: index only\nfallback: crawl if absent\nforce: request fresh crawl", TEAL),
                 ("Document", "text / html / markdown\nContent + returned status\nOptional freshness metadata", PURPLE)], y=146, h=177)
    box(s, "App selection", "fallback + renderDynamicPages=true\nmarkdown; requested maxLength 15,000", 36, 349, 423, 83, BLUE, 12.5)
    box(s, "Material limits", "202 means pending, not inspected.\nNo viewport PNG, clicks or freshness guarantee.", 491, 349, 433, 83, AMBER, 12.5)
    band(s, "Dynamic rendering support does not turn the Browse API response into an interactive browser session.")
    note(s, "Browse returns a structured document rather than the original page. Current docs allow text/html/markdown, "
         "not passage. liveCrawl defaults none; app chooses fallback, dynamic rendering, markdown and maxLength15000. "
         "force requests can return stale content; liveCrawlSucceeded appears as false for that case, never true. "
         "Do not infer freshness from an omitted flag or missing date. No viewport image field or interaction loop "
         "is provided by this response.", ("W3", "W4", "APP"))

    s = page(p, "Reading the data", "Search and Browse answer different input shapes",
             "Synthetic fragments below teach response structure; no publisher grounding is reproduced.",
             ("W1", "W2", "W3"))
    box(s, "SEARCH: query → many candidates", "", 36, 139, 427, 272, BLUE)
    text(s, '{ "webResults": [\n    { "url": "https://example.org/a",\n      "title": "<source title>",\n      "content": "<query passage>",\n      "crawledAt": "<optional date>" }\n  ]\n}', 53, 196, 394, 197, 12.3, NAVY, mono=True)
    box(s, "BROWSE: URL → one document", "", 495, 139, 429, 272, TEAL)
    text(s, '{ "url": "https://example.org/a",\n  "title": "<source title>",\n  "content": "<structured document>",\n  "crawledAt": "<optional date>",\n  "traceId": "<request identifier>"\n}', 512, 196, 395, 197, 12.3, NAVY, mono=True)
    band(s, "Neither a source URL nor a successful response proves that a claim is true. Preserve access and timing context.")

    s = page(p, "The browser stack", "Playwright ≠ CDP ≠ extension ≠ VM",
             "Understand the layers before comparing “browser” paths.",
             ("P5", "C1", "C2", "A1", "APP"))
    for i, (title, body, accent) in enumerate([
        ("Playwright", "Automation library: navigation, DOM, network events and screenshots.", BLUE),
        ("CDP", "Chrome DevTools Protocol: attach to an existing Chromium browser.", TEAL),
        ("Chrome extension", "Browser-installed code: messaging, script injection and tab capture.", PURPLE),
        ("VM host", "A virtual computer: runs an OS, worker, browser and dependencies.", AMBER),
    ]):
        y = 139 + i * 72
        rect(s, 36, y, 184, 58, accent)
        text(s, title, 50, y + 17, 158, 29, 16, WHITE, True)
        text(s, body, 247, y + 17, 669, 36, 14)
    band(s, "CDP is Chromium-only and lower-fidelity than Playwright's own protocol. A VM is placement, not an intelligence model.")

    s = page(p, "What is Azure BAT?", "Provision a cloud session; observe with Playwright",
             "The implemented BAT path is not a generic screenshot-to-click Computer Use model.",
             ("F1", "F3", "P5", "APP"))
    sequence(s, ["App + identity", "Foundry Toolbox", "Cloud Chromium"], [
        (0, 1, "browser_automation create_session", PURPLE),
        (1, 0, "Session handle + secure CDP endpoint", PURPLE),
        (0, 2, "Playwright connect_over_cdp", BLUE),
        (2, 0, "DOM / status / genuine viewport PNG", TEAL),
        (0, 1, "Session cleanup", AMBER),
    ])
    band(s, "Managed browser infrastructure helps operations; the page can still deny access, show a challenge or yield no useful evidence.")
    note(s, "Official BAT is GA, powered by Playwright Workspaces and exposed through tools/MCP. Vendor service "
         "includes broader navigation/form/live-view/take-control features. This app provisions a session via "
         "Foundry Toolbox browser_automation create_session and attaches Playwright through CDP for bounded "
         "read-only observation. No arbitrary user CDP endpoint, generic Computer Use model, live-view or human "
         "take-control capability is promised as a portal feature.", ("F1", "F3", "P5", "APP"))

    s = page(p, "What Playwright observes", "Three evidence channels can disagree",
             "An image of a loaded page is not the same as extracted product text.",
             ("P2", "P3", "APP"))
    box(s, "BROWSER PAGE", "Same navigation\nSame time window", 36, 227, 229, 139, NAVY)
    for y, title, body, accent in [
        (139, "Network", "Actual response status where observed;\nredirects, errors and request timing.", BLUE),
        (241, "Document", "DOM text / JSON-LD-derived fields;\nmay be empty or a verification page.", TEAL),
        (343, "Viewport PNG", "Pixels visible at capture time;\ncan depict a challenge, not the product.", PURPLE),
    ]:
        link(s, 270, 295, 352, y + 40, accent)
        box(s, title, body, 364, y, 560, 86, accent, 13)
    band(s, "Playwright also has network interception capabilities; those are not permission to alter evidence or evade a website.")

    s = page(p, "Contexts and placement", "A browser context is a session boundary",
             "Cookies, storage, locale and network location affect what a page serves.",
             ("P4", "P5", "APP"))
    for x, title, accent in [(36, "Cloud BAT", PURPLE), (343, "Local browser", TEAL), (650, "Azure VM", AMBER)]:
        rect(s, x, 149, 274, 216, PALE, BORDER)
        text(s, title, x + 17, 164, 240, 33, 19, accent, True)
        rect(s, x + 18, 214, 238, 123, WHITE, BORDER)
        text(s, "Browser context\nCookies + storage + pages\nIts own host / network", x + 32, 234, 208, 88, 14)
    text(s, "Isolated contexts help reproducibility; they do not make independent environments identical.",
         36, 391, 888, 34, 16, NAVY, True)
    band(s, "No method automatically inherits the buyer's login, permissions or local session. Public review access is not guaranteed.")

    edition.retain(16, ("P3", "C1", "C3", "APP", "EDIT"),
                   "Preserved the key implementation distinction: the interactive local portal path is a Chrome "
                   "extension bridge, while matched CLI benchmarks use isolated native Chrome through Playwright. "
                   "Do not treat benchmark success as validation of the visitor extension installation or session.")

    s = page(p, "Local extension / execution", "Messages coordinate; permissions enable observation",
             "Chrome APIs run in the visitor's browser, not inside the backend or BAT.",
             ("C1", "C2", "C3", "APP"))
    sequence(s, ["Portal page", "Extension worker", "Local tab"], [
        (0, 1, "Authenticated command → message bridge", BLUE),
        (1, 2, "Navigate + scripting.executeScript", TEAL),
        (2, 1, "DOM-derived observation", TEAL),
        (1, 2, "captureVisibleTab: active viewport", PURPLE),
        (1, 0, "Observation + permitted PNG", BLUE),
    ])
    band(s, "Review tabs / scripting / host permissions. Client-reported loaded=200 is not a measured origin HTTP response.")
    note(s, "Chrome messaging connects content scripts, service workers and extension contexts. scripting needs "
         "scripting plus host permission or temporary activeTab grant. captureVisibleTab captures the active tab's "
         "visible area and requires all_urls or activeTab under documented rules. The app uses tabs, activeTab, "
         "scripting and host access; do not claim all permissions come from activeTab alone. App client load status "
         "200 is not an observed origin response code. Capture scope is viewport, not a whole-site archive.",
         ("C1", "C2", "C3", "APP"))

    s = page(p, "What is the VM path?", "An operator-owned Windows browser worker",
             "Azure supplies the virtual computer; this app supplies the queue, worker and observation contract.",
             ("A1", "P3", "APP"))
    sequence(s, ["Portal queue", "Windows worker", "VM Chrome"], [
        (1, 0, "Authenticated heartbeat + poll", AMBER),
        (0, 1, "Bounded public-URL inspection command", BLUE),
        (1, 2, "Playwright drives native Chrome", TEAL),
        (2, 1, "Page observation + viewport PNG", TEAL),
        (1, 0, "Return result / restriction / failure", BLUE),
    ])
    band(s, "VM power state is not worker readiness. Check a fresh authenticated heartbeat; operate patches, profile hygiene and cost.")
    note(s, "The configured worker uses a Windows scheduled task, native Chrome and a worker profile. Headless "
         "native Chrome does not mean a visible desktop window. Current worker uses viewport screenshots and "
         "bounded navigation; actual response status is recorded where available. The operator owns OS/browser "
         "patches, availability, credential rotation and compute cost. No live/private screenshots embedded.",
         ("A1", "P3", "APP"))

    s = page(p, "Identity is not access to every website", "Separate service identity from website permission",
             "Managed identity can obtain Entra tokens; authorization still has to be granted.",
             ("A2", "F3", "W4", "APP"))
    row_flow(s, [("Workload identity", "App / Azure resource\nobtains an Entra token\nwithout an embedded secret.", BLUE),
                 ("Service authorization", "Foundry or another service\nchecks granted access,\nroles and connection policy.", PURPLE),
                 ("Website boundary", "The destination still controls\npublic access, login,\nrate limits and restrictions.", AMBER)], y=154, h=182)
    text(s, "This prototype also has distinct request-scoped WebIQ credentials and an authenticated worker channel.",
         40, 367, 880, 54, 16)
    band(s, "Do not claim the worker secret is replaced by managed identity, or that Azure RBAC grants arbitrary web access.")

    s = page(p, "The fifth path", "Search-only reviews skips review URL visits",
             "It still starts from an inspected listing. Search evidence is deliberately not promoted to page observation.",
             ("W1", "F2", "APP"))
    row_flow(s, [("Listing", "Default: WebIQ Browse\nOr explicitly choose another\nseed inspector.", BLUE),
                 ("Search", "Plan queries; retrieve\nattributed passages or\nToolbox search evidence.", TEAL),
                 ("No review visits", "Mark intentional skips.\nNo review-page DOM,\nno review-page PNGs.", AMBER),
                 ("Brief", "Cite search evidence.\nExplain unvisited pages\nand verification gaps.", NAVY)], y=151, h=183)
    text(s, "Useful when search context is sufficient for an initial brief; weaker when detail requires reading the page.",
         40, 368, 880, 50, 16)
    band(s, "An intentional skip is not a failed inspection. A Browse seed + Toolbox discovery still requires a deliberate credential-compatible setup.")

    s = page(p, "Five routes at a glance", "Same research story; different observation paths",
             "Native swimlanes show where evidence comes from. Discovery remains an independent choice.")
    headers = [(255, "LISTING"), (465, "DISCOVERY"), (678, "REVIEW VISITS"), (867, "OUTPUT")]
    for x, label in headers:
        text(s, label, x - 74, 135, 148, 21, 10.5, GRAY, True, align=PP_ALIGN.CENTER)
    for i, (_, label, accent) in enumerate(METHODS):
        y = 174 + i * 48
        text(s, label, 36, y + 8, 158, 32, 12.5, accent, True)
        listing = ["Browse", "Cloud BAT", "Local", "VM worker", "Browse*"][i]
        reviews = ["Browse", "Cloud BAT", "Local", "VM worker", "SKIP"][i]
        for x, value in [(255, listing), (465, "WebIQ / Toolbox"), (678, reviews), (867, "Cited brief")]:
            rect(s, x - 74, y, 148, 35, PALE, BORDER)
            text(s, value, x - 68, y + 8, 136, 21, 11.5, accent if value == "SKIP" else INK,
                 value == "SKIP", align=PP_ALIGN.CENTER)
        for a, b in [(329, 387), (539, 600), (752, 789)]:
            link(s, a, y + 17, b, y + 17, accent)
    band(s, "* Search-only may explicitly select another seed inspector. Local portal extension and CLI runner are different implementations.")

    s = page(p, "Browser verification / current policy", "Wait for automatic clearance — never solve a puzzle",
             "The current browser inspector can monitor a transient verification screen in the same session.")
    box(s, "Observe page", "Read status + page signals", 36, 152, 246, 108, BLUE)
    box(s, "Hard denial / login / 429 / 451", "Skip immediately; no retries.", 332, 139, 592, 84, AMBER, 13)
    link(s, 287, 188, 327, 181, AMBER)
    box(s, "Possible automatic clearance", "Wait 3 seconds → recheck same session.\nAt most 2 rechecks, then record the outcome.", 332, 249, 592, 110, TEAL, 13)
    link(s, 287, 226, 327, 282, TEAL)
    text(s, "Clears: inspect normally.  Still restricted: skip with an explicit evidence gap.",
         336, 386, 580, 44, 14, NAVY, True)
    band(s, "No human pause, solver, proxy/stealth change or inspector switch. Interactive CAPTCHA puzzles are not solved.")
    note(s, "Current implementation distinguishes transient automatic verification from hard denials. Up to two "
         "same-session rechecks after three-second waits; no human pause, CAPTCHA solver, proxy, stealth changes "
         "or inspector switch. Hard denials, login requirements, 429 and 451 skip without these retries. This is "
         "bounded monitoring, not a universal CAPTCHA bypass or a guarantee of clearance.", ("APP",))

    s = page(p, "Public-page scope", "A platform is a destination, not a special exemption",
             "Reddit, Amazon and YouTube are not blanket-excluded; access and useful extraction are not guaranteed.")
    row_flow(s, [("Public text", "Listing, lab review,\nowner discussion, page text\nand visible comments.", BLUE),
                 ("Respect the boundary", "Login, denial, rate limit\nor persistent verification\nbecomes an explicit skip.", AMBER),
                 ("No media pipeline", "No video downloading,\ntranscription or Content\nUnderstanding processing.", PURPLE)], y=155, h=178)
    text(s, "A public YouTube review page can be visited for available page text/comments.\nThat is not watching, transcribing or analyzing the video's media.",
         40, 365, 880, 61, 16)
    band(s, "Source relevance does not imply access, credibility or permission to reuse its content.")

    edition.retain(24, ("W3", "P3", "APP", "EDIT"),
                   "Preserved the three evidence states: extracted document text, genuine browser PNG, and absent "
                   "evidence. A synthetic diagnostic is not a captured website. A challenge screenshot can be genuine "
                   "without yielding product evidence. SHA256 binds bytes, not factual truth or freshness.")

    s = page(p, "Annotated output / synthetic", "Read the observation as a receipt",
             "A normalized shape is useful for comparison; it is not an unlimited raw vendor payload.")
    rect(s, 36, 142, 510, 281, NAVY)
    text(s, '{\n  "url": "https://example.org/review",\n  "inspector": "vm_playwright",\n  "status": "restricted_by_site",\n  "fetched_characters": 0,\n  "capture_kind": "browser_viewport",\n  "sha256": "<hash of actual PNG bytes>"\n}', 54, 163, 475, 246, 14, "DCF3FF", mono=True)
    for y, title, body, accent in [
        (143, "Status", "Tells us what access actually allowed.", AMBER),
        (240, "Text length", "Zero is not hidden evidence to invent.", BLUE),
        (337, "Capture + hash", "Genuine pixels may show only a challenge.", PURPLE),
    ]:
        box(s, title, body, 578, y, 346, 86, accent, 12)
    band(s, "Synthetic teaching fragment only. No live publisher grounding, secrets or private screenshots appear in this deck.")

    s = page(p, "Sony walkthrough / claim to reality", "Follow one priority all the way to uncertainty",
             "Illustrative reasoning, not an asserted conclusion about the WH-1000XM6.")
    row_flow(s, [("Buyer priority", "Comfort with thin glasses\non a 10-hour flight.", BLUE),
                 ("Evidence needed", "Test duration + seal context;\nindependent measurements\nand contextual owner reports.", TEAL),
                 ("Safe interpretation", "Supported only under\nobserved conditions;\nthis buyer may need a trial.", PURPLE)], y=152, h=181)
    rect(s, 36, 362, 888, 64, PALE, BORDER)
    text(s, "For hinge durability: reported incidents are signals, not a measured failure rate.\nFor budget: no observed current offer means no invented qualifying price.", 51, 374, 858, 47, 14)
    band(s, "Every claim check needs a source URL, evidence type, buyer impact and a clearly stated limitation.")

    edition.retain(26, ("APP", "EDIT"), "Preserved the structured synthesis contract and explicit prompt rules. "
                   "GPT-5.6-Terra is the configured app deployment. No unverified public model specifications are "
                   "asserted. A schema constrains response structure; it does not establish truth. Timeouts and "
                   "exact budgets are implementation details, not service SLAs.")

    s = page(p, "Evidence interpretation", "Do not collapse different support into “verified”",
             "The type and limits of evidence matter more than a generic completion badge.")
    for i, (label, detail, accent) in enumerate([
        ("Seller claim", "The listing says it; independent support is still needed.", BLUE),
        ("Measured result", "Check conditions, equipment, variant and test scope.", TEAL),
        ("Owner report", "Useful lived context; anecdotes do not establish prevalence.", PURPLE),
        ("Model inference", "Reasoned interpretation, explicitly separated from observation.", NAVY),
        ("Evidence gap", "Unavailable / restricted / unvisited / outdated / not corroborated.", AMBER),
    ]):
        y = 139 + i * 57
        rect(s, 36, y, 190, 44, accent)
        text(s, label, 50, y + 12, 161, 27, 13.5, WHITE, True)
        text(s, detail, 250, y + 12, 674, 34, 14)
    band(s, "Conflicting sources should remain visible. “Completed synthesis” is not “all claims verified.”")

    edition.retain(27, ("APP", "RUN", "EDIT"), "Preserved editable replay schematic. Replay reads permitted stored "
                   "operational events and separate browser artifacts, not raw WebIQ grounding. Associated "
                   "attribution/generated brief may remain briefly in RAM. Expired context must stay expired; "
                   "replay itself never invokes the research model or service endpoints and does not require "
                   "an online VM. Post-round VM deallocation is separate from the recorded execution timing.")

    s = page(p, "Evidence lifecycle", "Replay is not a publisher-content archive",
             "Separate transient grounding, short-lived output and permitted browser observations.",
             ("W5", "APP"))
    for y, title, during, after, accent in [
        (143, "Raw WebIQ grounding", "Workflow-only model input", "DELETE at workflow completion", AMBER),
        (240, "Generated brief + attribution", "User-facing cited synthesis", "Briefly in RAM; then expires", PURPLE),
        (337, "Operational / browser evidence", "Status, timing, separate observations", "Permitted metadata + hash-verified PNGs", BLUE),
    ]:
        text(s, title, 36, y + 12, 218, 57, 13.8, accent, True)
        box(s, during, "", 270, y, 287, 75, accent)
        link(s, 563, y + 38, 597, y + 38, accent)
        box(s, after, "", 609, y, 315, 75, accent)
    band(s, "Safe expanded panels remain projections/excerpts. No raw WebIQ reconstruction or export from expired context.")

    if runs:
        add_results(p, data, runs, runs_path, not_before)

    s = page(p, "Choosing a route", "Choose the lightest evidence path that fits the question",
             "This is a capability decision tree, not a measured quality or cost ranking.",
             ("W3", "F1", "P3", "A1", "APP"))
    box(s, "Need review-page details?", "No → search-only initial brief\nYes → inspect eligible review URLs", 36, 145, 374, 115, BLUE)
    box(s, "Need rendered state / actual PNG?", "No → WebIQ Browse may be sufficient\nYes → choose a browser placement", 36, 297, 374, 115, TEAL)
    link(s, 224, 264, 224, 292)
    for y, title, body, accent in [
        (139, "Managed cloud", "BAT: managed browser session + lifecycle", PURPLE),
        (240, "Visitor's browser", "Local extension: install + connect + review permissions", TEAL),
        (341, "Operator-owned host", "VM worker: availability + maintenance + credentials", AMBER),
    ]:
        box(s, title, body, 492, y, 432, 86, accent, 12.5)
    link(s, 416, 355, 480, 281, BLUE)
    band(s, "All routes can fail. Evidence gaps and operating ownership belong in the decision.")

    s = page(p, "Human action boundary", "Research does not authorize commerce",
             "The cart branch is a separate simulation, not retailer mutation or payment.")
    row_flow(s, [("Cited brief", "Buyer reads sources,\ntrade-offs and unknowns.", BLUE),
                 ("Optional approval", "Explicit human decision\nfor the demonstration branch.", PURPLE),
                 ("Cart SIMULATION", "Show intended action\nwithout mutating a real cart.", TEAL)], y=153, h=176)
    rect(s, 36, 363, 888, 61, NAVY)
    text(s, "NO CHECKOUT   •   NO PURCHASE   •   NO PAYMENT   •   NO RETAILER CART MUTATION",
         53, 383, 854, 30, 15, WHITE, True, align=PP_ALIGN.CENTER)
    band(s, "A future commerce integration would require a new explicit authorization, identity and risk design.")

    s = page(p, "Shipped vs next revamp", "Do not teach a roadmap item as working behavior",
             "Keep current policy, available vendor capabilities and future app ideas separate.")
    box(s, "SHIPPED NOW", "Public listing + search + optional review visits\nSame selected inspector across visited stages\nBounded automatic-clearance monitoring\nCited synthesis + simulation-only action", 36, 145, 427, 267, TEAL, 14)
    box(s, "NEXT REVAMP — NOT SHIPPED", "Targeted search-guided navigation\nto relevant comment sections\n\nWould need explicit navigation policies,\nprovenance, extraction checks and validation.", 495, 145, 429, 267, PURPLE, 14)
    band(s, "No planned comment-section navigator, media processing, CAPTCHA solving or checkout is implied by the current demo.")

    edition.retain(30, ("APP", "RUN", "EDIT"), "Preserved the user's discussion framework: evidence quality, "
                   "execution placement and human accountability. Operational five-path receipts inform execution "
                   "discussion; they do not establish recommendation quality, source credibility or a universal winner.")
    add_sources(p)
    check = validate_package(p)
    if hashlib.sha256(BASE.read_bytes()).hexdigest() != edition.base_hash:
        raise RuntimeError("Source presentation changed during generation; rerun against its new edited version.")
    if output.resolve() == BASE.resolve():
        raise ValueError("Refusing to overwrite the user's edited presentation.")
    output.parent.mkdir(parents=True, exist_ok=True)
    p.save(output)
    metrics = {
        "output": str(output), "slides": len(p.slides), "source_sha256": edition.base_hash,
        "retained_user_slides": edition.inherited,
        "official_sources": len([k for k in SOURCES if k not in {"APP", "RUN", "EDIT"}]),
        "native_shapes": sum(len(s.shapes) for s in p.slides),
        "native_diagram_shapes": sum(sh.name.startswith("diagram-") for s in p.slides for sh in s.shapes),
        "connector_count": sum(sh.shape_type == 9 for s in p.slides for sh in s.shapes),
        "tables": sum(sh.has_table for s in p.slides for sh in s.shapes),
        "embedded_pictures": sum(sh.shape_type == 13 for s in p.slides for sh in s.shapes),
        "slides_with_diagrams": sum(any(sh.name.startswith("diagram-") for sh in list(s.shapes)[1:]) for s in p.slides),
        "research_only": research_only, "validation": check,
        "run_ids": {m: r["run_id"] for m, r in (runs or {}).items()},
    }
    return metrics


def add_results(p, data, runs, path, not_before):
    example = runs["webiq_browse"]
    context = (
        f"Selected latest receipt per path at/after {not_before.isoformat()}. "
        f"Revision {example['server_revision']}; source {example['source_commit']}; "
        f"Round {data.get('round_id', 'not specified')}; JSON recorded_at={data.get('recorded_at')}; "
        f"SHA256={hashlib.sha256(path.read_bytes()).hexdigest()}. "
        "Same product URL, intent, primary query, model deployment and runtime revision. Different discovered "
        "source sets, browsers, networks and live access conditions remain uncontrolled. This is operational "
        "comparison only, not a controlled recommendation-quality experiment, cost ranking or service SLA. "
        "Execution revision is taken from each run receipt, never from the later report-publication deployment. "
        "Review inspection counts exclude the seed listing; PNG counts include the whole run. "
        + "\n".join(f"{label}: {runs[m]['run_id']} | {runs[m]['started_at']} | "
                    f"outcome={runs[m].get('outcome')} | profile={runs[m].get('profile')}"
                    for m, label, _ in METHODS)
    )
    s = page(p, "Fresh five-path round / measured", "Elapsed time is operational, not answer quality",
             f"Sony case | Round {data.get('round_id', 'not specified')} | Execution {example['server_revision']}",
             ("RUN",))
    peak = max(float(r["elapsed_s"]) for r in runs.values()) or 1
    for i, (method, label, accent) in enumerate(METHODS):
        r = runs[method]
        y = 145 + i * 56
        text(s, label, 36, y + 8, 188, 35, 13.5, accent, True)
        rect(s, 237, y + 3, 500, 29, PALE, kind=MSO_SHAPE.RECTANGLE)
        rect(s, 237, y + 3, max(2, float(r["elapsed_s"]) / peak * 500), 29, accent,
             kind=MSO_SHAPE.RECTANGLE)
        text(s, f"{r['elapsed_s']:.2f}s", 750, y + 6, 108, 28, 14, accent, True)
        text(s, str(r.get("outcome", "unknown")), 852, y + 8, 72, 29, 11, GRAY)
        start = next((e["elapsed_s"] for e in r.get("events", [])
                      if e["type"] == "status" and e.get("stage") == "permitted_source_browsing"), None)
        finish = next((e["elapsed_s"] for e in r.get("events", [])
                       if e["type"] == "status" and e.get("stage") == "synthesize_verdict"), None)
        detail = ("Review visits intentionally skipped" if method == "search_only"
                  else f"Review stage: {finish - start:.2f}s" if start is not None and finish is not None
                  else "Review stage: timing not recorded")
        text(s, detail, 237, y + 35, 650, 18, 10, GRAY)
    band(s, "Matching inputs/runtime do not equal matching sources or access conditions. One run per path cannot establish a universal winner.")
    note(s, context + "\nReview-stage time is the recorded interval from permitted_source_browsing status to "
         "synthesize_verdict status, including access checks and transition overhead; it is not isolated provider latency.",
         ("RUN",))

    s = page(p, "Fresh five-path round / coverage", "Coverage and screenshots answer different questions",
             "Inspected / discovered counts review URLs only. PNG counts include listing and other whole-run captures.",
             ("RUN",))
    for x, label in [(36, "PATH"), (258, "REVIEW COVERAGE"), (590, "GENUINE PNGs"), (756, "SYNTHESIS")]:
        text(s, label, x, 136, 170, 22, 10, GRAY, True)
    for i, (method, label, accent) in enumerate(METHODS):
        r = runs[method]
        y = 173 + i * 48
        total, done = int(r["discovered"]), int(r["inspected"])
        text(s, label, 36, y + 5, 199, 33, 13, accent, True)
        rect(s, 258, y + 5, 200, 23, PALE, kind=MSO_SHAPE.RECTANGLE)
        if total and method != "search_only":
            rect(s, 258, y + 5, max(1, min(done / total, 1) * 200), 23, accent,
                 kind=MSO_SHAPE.RECTANGLE)
        text(s, f"{done}/{total}" if method != "search_only" else f"0/{total} — skipped", 470, y + 5, 119, 27, 12)
        text(s, str(r["real_captures"]), 631, y + 3, 75, 29, 17, accent, True)
        text(s, str(r.get("synthesis_status", "unknown")), 756, y + 5, 168, 28, 12)
    total_pngs = sum(r["real_captures"] for r in runs.values())
    bat_sources = runs["live"].get("inspections", [])
    rechecks = sum(v.get("verification_attempts", 0) for v in bat_sources)
    unresolved = sum(v.get("verification_outcome") == "unresolved" for v in bat_sources)
    text(s, f"{total_pngs} real PNG capture events, not unique pages. BAT: {rechecks} rechecks across "
         f"{unresolved} unresolved review URLs.", 36, 416, 888, 26, 12.5, NAVY, True)
    band(s, "Search-only skips are intentional. Zero Browse PNGs is expected. A completed synthesis can still have material evidence gaps.")
    note(s, context + f"\n{total_pngs} genuine whole-round capture events. BAT final per-URL receipts record "
         f"{rechecks} automatic-verification rechecks across {unresolved} unresolved review pages. These rechecks "
         "do not imply interactive CAPTCHA solving. The VM can be deallocated after measurement; replay "
         "does not require an online worker.", ("RUN",))

    s = page(p, "Fresh five-path round / context volume", "More context does not mean more truth",
             "Gross character totals from bounded search and inspected-review context; not tokens or net-new verified facts.",
             ("RUN",))
    channels = (("search_content_characters", "Search passages", TEAL),
                ("inspected_context_characters", "Inspected-page context", PURPLE))
    for x, (_, label, accent) in zip((260, 575), channels):
        rect(s, x, 136, 20, 11, accent, kind=MSO_SHAPE.RECTANGLE)
        text(s, label, x + 29, 133, 265, 22, 11.5, accent, True)
    peak = max((r.get("evidence_metrics", {}).get(key, 0)
                for r in runs.values() for key, _, _ in channels), default=1) or 1
    for i, (method, label, accent) in enumerate(METHODS):
        r = runs[method]
        y = 174 + i * 49
        text(s, label, 36, y + 4, 195, 32, 13, accent, True)
        for j, (key, _, channel_color) in enumerate(channels):
            value = r.get("evidence_metrics", {}).get(key)
            yy = y + j * 17
            rect(s, 260, yy, 460, 12, PALE, kind=MSO_SHAPE.RECTANGLE)
            if value:
                rect(s, 260, yy, max(1, value / peak * 460), 12, channel_color,
                     kind=MSO_SHAPE.RECTANGLE)
            text(s, f"{value:,}" if value is not None else "not recorded", 739, yy - 3,
                 185, 20, 10.5, channel_color, True)
    band(s, "Search and page text can overlap or conflict. Additional characters do not prove authenticity, corroboration or recommendation accuracy.")
    note(s, context + "\nMetrics are sum(len(source.content[:3000])) for discovery and "
         "sum(len(source.page_content[:6000] or source.observed_excerpt)) for successfully inspected review "
         "sources. They exclude seed-listing context and are not deduplicated for semantic overlap. "
         "Missing metrics are labeled not recorded, never assumed to be zero. No actual grounding is displayed.",
         ("RUN", "APP"))

    s = page(p, "Fresh five-path round / source sets", "Same query does not force the same source mix",
             "Candidate-set overlap explains why matched inputs are not a controlled recommendation-quality experiment.",
             ("RUN",))
    url_sets = [set(candidate_urls(runs[method])) for method, _, _ in METHODS]
    common, distinct = len(set.intersection(*url_sets)), len(set.union(*url_sets))
    domain_counts = {}
    for i, (method, label, accent) in enumerate(METHODS):
        r = runs[method]
        domains = Counter(urlsplit(url).netloc.removeprefix("www.") for url in candidate_urls(r))
        domain_counts[method] = len(domains)
        y = 147 + i * 55
        rect(s, 36, y + 9, 15, 15, accent, kind=MSO_SHAPE.OVAL)
        text(s, label, 65, y, 213, 27, 13.5, accent, True)
        text(s, f"{len(domains)} discovered domains", 65, y + 27, 213, 23, 11.5, GRAY)
        link(s, 283, y + 24, 343, 219, accent)
    box(s, f"{common} shared exact URLs", "The intersection:\ncandidates discovered\nby every path.", 350, 148, 270, 150, BLUE)
    link(s, 626, 219, 674, 219, BLUE)
    text(s, "subset of", 627, 193, 55, 19, 9.5, GRAY)
    box(s, f"{distinct} distinct URLs", "The union:\nall candidates discovered\nacross this round.", 685, 148, 239, 150, TEAL)
    box(s, "What is still uncontrolled?", "Secondary queries, candidate source sets, browser/network conditions\n"
        "and site access can differ. Failed or intentionally unvisited\ncandidates still belong to the discovery set.",
        350, 327, 574, 106, PURPLE, 12.5)
    band(s, "Presence is not endorsement. Candidate counts do not score credibility, corroboration or the quality of the recommendation.")
    note(s, context + "\nDomains are derived from enrichment-event URL metadata, with inspection URLs only as a "
         "fallback. This includes deliberately unvisited search-only candidates, not retained search grounding. "
         f"Exact URL intersection={common}; union={distinct}. "
         f"Distinct discovered-domain counts: {domain_counts}. "
         "If metadata is absent the slide reports that absence; it never invents a domain or source.", ("RUN",))


def candidate_urls(run):
    for event in reversed(run.get("events", [])):
        if event.get("type") == "enrichment" and event.get("urls"):
            return list(dict.fromkeys(url for url in event["urls"] if isinstance(url, str) and url))
    return list(dict.fromkeys(v["url"] for v in run.get("inspections", []) if v.get("url")))


def add_sources(p):
    keys = [k for k in SOURCES if k not in {"APP", "RUN", "EDIT"}]
    titles = (
        "WebIQ content and cloud browser foundations",
        "Foundry search, Toolbox and Playwright observation",
        "CDP, Chrome extensions and Azure execution",
        "LangGraph and progressive-disclosure skills",
    )
    for offset in range(0, len(keys), 6):
        subset = keys[offset:offset + 6]
        s = page(p, "Official sources / verified " + VERIFIED, titles[offset // 6],
                 "Slide notes connect each capability to its source. Links below are clickable.", tuple(subset))
        for i, key in enumerate(subset):
            title, url = SOURCES[key]
            y = 136 + i * 48
            text(s, key, 36, y + 3, 40, 29, 13, BLUE, True)
            shape = text(s, title, 84, y, 840, 25, 13, NAVY, True)
            shape.text_frame.paragraphs[0].runs[0].hyperlink.address = url
            shown = url.replace("https://", "")
            shape = text(s, shown, 84, y + 24, 840, 21, 9.1, GRAY)
            shape.text_frame.paragraphs[0].runs[0].hyperlink.address = url
        band(s, "Capabilities change. Official service documentation, configured app behavior and measured receipts are different evidence layers.")
        note(s, "All listed official pages were directly retrieved and checked on " + VERIFIED
             + ". The CDP source document corroborates Chromium-only / lower-fidelity caveats. "
             "No marketing benchmark, unverified public model specification or third-party illustration is reused.", tuple(subset))


def validate_package(prs):
    slide_ids = [sid.id for sid in prs.slides._sldIdLst]
    if len(slide_ids) != len(set(slide_ids)):
        raise ValueError("Slide identifiers must remain unique when retaining edited source slides.")
    boundaries = []
    tables = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.left < -Pt(0.5) or sh.top < -Pt(0.5) or sh.left + sh.width > prs.slide_width + Pt(0.5) or sh.top + sh.height > prs.slide_height + Pt(0.5):
                boundaries.append((i, sh.name))
            if sh.has_table:
                tables.append(i)
        if not s.has_notes_slide or "Sources" not in s.notes_slide.notes_text_frame.text and "[" not in s.notes_slide.notes_text_frame.text:
            raise ValueError(f"Slide {i} has no source notes")
    if boundaries or tables:
        raise ValueError(f"Shape boundaries={boundaries}; unexpected tables on slides={tables}")
    return {"shape_boundary_errors": boundaries, "table_slides": tables}


def render_and_audit(output, directory):
    """Use existing desktop PowerPoint without quitting or touching another presentation."""
    import win32com.client

    directory.mkdir(parents=True, exist_ok=True)
    app = win32com.client.Dispatch("PowerPoint.Application")
    opened = None
    audit = {"slides": [], "overflow": [], "shape_boundaries": [], "renderer": "Microsoft PowerPoint COM"}
    try:
        opened = app.Presentations.Open(str(output.resolve()), ReadOnly=True, Untitled=False, WithWindow=False)
        width, height = float(opened.PageSetup.SlideWidth), float(opened.PageSetup.SlideHeight)
        for slide in opened.Slides:
            slide.Export(str(directory / f"slide-{slide.SlideIndex:02}.png"), "PNG", 1600, 900)
            audit["slides"].append(slide.SlideIndex)
            for sh in slide.Shapes:
                if sh.Left < -0.5 or sh.Top < -0.5 or sh.Left + sh.Width > width + 0.5 or sh.Top + sh.Height > height + 0.5:
                    audit["shape_boundaries"].append({"slide": slide.SlideIndex, "shape": sh.Name})
                if sh.HasTextFrame and sh.TextFrame.HasText:
                    tf = sh.TextFrame2
                    if (tf.TextRange.BoundHeight > sh.Height - tf.MarginTop - tf.MarginBottom + 2
                            or tf.TextRange.BoundWidth > sh.Width - tf.MarginLeft - tf.MarginRight + 2):
                        audit["overflow"].append({"slide": slide.SlideIndex, "shape": sh.Name,
                                                  "bound_height": round(tf.TextRange.BoundHeight, 2),
                                                  "height": round(sh.Height, 2),
                                                  "bound_width": round(tf.TextRange.BoundWidth, 2),
                                                  "width": round(sh.Width, 2),
                                                  "text": sh.TextFrame.TextRange.Text[:110]})
        (directory / "powerpoint-audit.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    finally:
        if opened is not None:
            opened.Close()
        # Never quit a shared PowerPoint instance or close unrelated presentations.
    if audit["overflow"] or audit["shape_boundaries"]:
        raise ValueError(f"PowerPoint layout audit failed: {json.dumps(audit['overflow'] + audit['shape_boundaries'])}")
    return {"rendered_slides": len(audit["slides"]), "overflow": 0, "shape_boundary_errors": 0}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--runs", type=Path, default=RUNS)
    parser.add_argument("--run-not-before", default="2026-10-21T01:00:00Z",
                        help="Fresh-round lower timestamp bound; historical mixed runs are refused.")
    parser.add_argument("--research-only", action="store_true")
    parser.add_argument("--render-dir", type=Path)
    args = parser.parse_args()
    cutoff = iso_time(args.run_not_before)
    if cutoff.tzinfo is None:
        parser.error("--run-not-before must include a timezone.")
    result = build(args.output, args.runs, cutoff, args.research_only)
    if args.render_dir:
        result["render"] = render_and_audit(args.output, args.render_dir)
        (args.render_dir / "deck-metrics.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
