"""Optional authoring checks use synthetic test assets, never paid live workflows."""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace
import zipfile

import pytest

pytest.importorskip("pptx")
pytest.importorskip("PIL")
from PIL import Image
from pptx import Presentation
from pptx.util import Pt

from scripts import build_portal_walkthrough_deck as deck
from scripts import capture_portal_walkthrough as capture


def fixture_receipt(tmp_path):
    content = io.BytesIO()
    Image.new("RGB", (1000, 400), "white").save(content, "PNG")
    raw = content.getvalue()
    keys = ["01-home", *[key for key, *_ in deck.STEPS], "14-export"]
    shots = []
    for key in keys:
        (tmp_path / (key + ".png")).write_bytes(raw)
        shots.append({"key": key, "file": key + ".png", "sha256": hashlib.sha256(raw).hexdigest(),
                      "kind": "genuine_portal_screenshot", "content_modified": False, "width": 1000,
                      "height": 400, "captured_at": "2026-10-21T00:00:00Z",
                      "page_url": "https://portal.example/", "region": "synthetic test fixture"})
    data = {
        "schema": "portal-walkthrough-v1", "test_fixture": True, "finished_at": "2026-10-21T00:01:00Z",
        "started_at": "2026-10-21T00:00:00Z", "paid_runs_requested": 1, "raw_grounding_saved": False,
        "screens": shots, "product_url": deck.PRODUCT, "user_intent": "Synthetic layout test only",
        "revision": "test-fixture", "saved_research_url": "https://portal.example/saved-research?run=test",
        "saved_run": {"run_id": "fixture", "genuine_screenshots": 0, "claim_checks": 0,
                      "community_actions": 0, "executive_summary": "Synthetic layout fixture, not product findings."},
        "ui_outcome": "Synthetic test", "export_verified": {"genuine_pngs": 0},
    }
    path = tmp_path / "walkthrough.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def source_deck(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    labels = {2: "BUYER-FIRST FRAMING", 6: "SONY WALKTHROUGH / LISTING",
              7: "SONY WALKTHROUGH / QUERY PLANNING"}
    for index in range(1, 8):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Pt(36), Pt(30), Pt(888), Pt(60))
        box.text = labels.get(index, f"Preserved user content {index}")
    path = tmp_path / "base.pptx"
    prs.save(path)
    return path


def test_requires_explicit_live_opt_in_before_reading_credentials(tmp_path):
    args = SimpleNamespace(portal_url="https://portal.example", output=tmp_path / "new", run_live=False)
    with pytest.raises(ValueError, match="--run-live"):
        capture.capture_run(args)


def test_existing_capture_refuses_a_repeat_run(tmp_path):
    (tmp_path / "existing.png").write_bytes(b"fixture")
    args = SimpleNamespace(portal_url="https://portal.example", output=tmp_path, run_live=True)
    with pytest.raises(FileExistsError, match="must not be repeated"):
        capture.capture_run(args)


@pytest.mark.parametrize("paid_runs,origin", [(0, "https://portal.example"), (1, "https://wrong.example")])
def test_saved_capture_rejects_a_different_or_unsubmitted_run_before_login(tmp_path, paid_runs, origin):
    (tmp_path / "walkthrough.json").write_text(
        json.dumps({"paid_runs_requested": paid_runs, "portal_url": origin}), encoding="utf-8")
    args = SimpleNamespace(output=tmp_path, portal_url="https://portal.example")
    with pytest.raises(ValueError, match="already-submitted"):
        capture.resume_saved(args)


def test_focused_screenshot_records_crop_without_changing_pixels(tmp_path):
    buffer = io.BytesIO()
    Image.new("RGB", (800, 900), (23, 45, 67)).save(buffer, "PNG")

    class Element:
        def scroll_into_view_if_needed(self):
            pass

        def screenshot(self, **kwargs):
            return buffer.getvalue()

    shot = capture.capture(SimpleNamespace(url="https://portal.example/saved-research"),
                           tmp_path, "focused", "#observations", max_height=300, locator=Element())
    assert shot["region_crop"] == {"left": 0, "top": 0, "width": 800, "height": 300}
    assert shot["content_modified"] is False
    with Image.open(tmp_path / shot["file"]) as image:
        assert image.size == (800, 300)
        assert image.getpixel((799, 299)) == (23, 45, 67)


@pytest.mark.parametrize("url", ["http://portal.example", "https://portal.example/path", "https://portal.example?q=x"])
def test_origin_validation(url):
    with pytest.raises(ValueError):
        capture.validate_origin(url)


def test_receipt_hash_mismatch_is_rejected(tmp_path):
    path = fixture_receipt(tmp_path)
    (tmp_path / "01-home.png").write_bytes(b"not the captured image")
    with pytest.raises(ValueError, match="PNG|hash mismatch"):
        deck.load_receipt(path)


def test_synthetic_layout_build_preserves_input_and_slide_identifiers(tmp_path):
    path = fixture_receipt(tmp_path)
    base = source_deck(tmp_path)
    before = hashlib.sha256(base.read_bytes()).hexdigest()
    receipt_before = path.read_bytes()
    output = tmp_path / "walkthrough.pptx"
    summary = deck.build(base, path, output)
    assert hashlib.sha256(base.read_bytes()).hexdigest() == before
    assert path.read_bytes() == receipt_before
    assert summary["walkthrough_slides"] == 17
    assert summary["replaced_product_framing_slides"] == [2, 6, 7]
    assert len(summary["retained_slides"]) == 4
    with zipfile.ZipFile(output) as package:
        assert package.testzip() is None
    result = Presentation(output)
    identifiers = [sid.id for sid in result.slides._sldIdLst]
    assert len(identifiers) == len(set(identifiers))
    texts = "\n".join(shape.text for slide in result.slides for shape in slide.shapes if shape.has_text_frame)
    assert "Preserved user content 1" in texts
    assert "SONY WALKTHROUGH / LISTING" not in texts
    assert "Synthetic test" in texts
    assert sum(shape.shape_type == 13 for slide in result.slides for shape in slide.shapes) == len(json.loads(path.read_text())["screens"])
    assert any(shape.click_action.hyperlink.address == deck.PRODUCT for shape in result.slides[1].shapes)
    with pytest.raises(ValueError, match="never overwritten"):
        deck.build(base, path, output)


def test_provenance_notes_survive_an_edited_empty_notes_master(tmp_path):
    receipt = fixture_receipt(tmp_path)
    base = source_deck(tmp_path)
    prs = Presentation(base)
    for placeholder in list(prs.notes_master.placeholders):
        placeholder.element.getparent().remove(placeholder.element)
    prs.save(base)
    output = tmp_path / "with-notes.pptx"
    deck.build(base, receipt, output)
    assert "Live walkthrough provenance" in Presentation(output).slides[0].notes_slide.notes_text_frame.text
