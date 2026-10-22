"""Offline checks for the editable, product-neutral concept diagrams."""
import hashlib
import zipfile

import pytest

pytest.importorskip("pptx")
pytest.importorskip("PIL")
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from scripts import build_solution_concept_deck as concept


def test_builds_five_editable_diagrams_with_only_attributed_icon_pictures(tmp_path):
    output = tmp_path / "concept.pptx"
    result = concept.build(output)
    prs = Presentation(output)
    assert result["slides"] == 5
    assert result["provider_calls"] == 0
    assert result["output_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert (prs.slide_width, prs.slide_height) == (Pt(960), Pt(540))
    assert result["diagram_structure_native_editable"] is True
    assert result["icon_provenance"]["embedded_occurrences"] == 14
    for slide, expected_icons in zip(prs.slides, [5, 0, 3, 4, 2]):
        assert sum(shape.shape_type == MSO_SHAPE_TYPE.LINE for shape in slide.shapes) >= 3
        assert not any(shape.has_table for shape in slide.shapes)
        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pictures) == expected_icons
        for picture in pictures:
            assert picture.name.startswith("azure-icon-")
            key = picture.name.removeprefix("azure-icon-")
            assert picture.image.blob == (concept.ICON_DIR / (key + ".png")).read_bytes()
            assert picture.width == picture.height
        assert "Implementation inspected" in slide.notes_slide.notes_text_frame.text
        for shape in slide.shapes:
            assert min(shape.left, shape.top) >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
    with zipfile.ZipFile(output) as package:
        assert package.testzip() is None
        assert len(package.namelist()) == len(set(package.namelist()))


def test_diagrams_preserve_actual_runtime_and_evidence_distinctions(tmp_path):
    output = tmp_path / "concept.pptx"
    concept.build(output)
    prs = Presentation(output)
    texts = ["\n".join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in prs.slides]
    assert "FastAPI + SSE portal" in texts[0]
    assert "WebIQ Search OR Toolbox web_search" in texts[0]
    assert "Search-only reviews: skip stage 3" in texts[1]
    assert "never checkout" in texts[1]
    assert "Reopening saved research is a read" in texts[2]
    assert "Invalid-key errors do not silently switch" in texts[3]
    assert "Raw grounding is NOT archived" in texts[4]
    assert "not per-user isolation" in texts[4]
    assert all("Surface Laptop" not in value and "Sony" not in value for value in texts)
    assert "Foundry Models" in texts[0]
    assert "AZURE CONTAINER APPS ENVIRONMENT" in texts[0]
    assert "Azure Storage / Blob" in texts[0]


def test_icon_validation_rejects_a_modified_source_asset(tmp_path, monkeypatch):
    import shutil
    copied = tmp_path / "icons"
    shutil.copytree(concept.ICON_DIR, copied)
    (copied / "foundry-models.svg").write_bytes(b"modified icon")
    monkeypatch.setattr(concept, "ICON_DIR", copied)
    with pytest.raises(ValueError, match="hash mismatch"):
        concept.validate_icon_assets()


def test_does_not_overwrite_an_existing_presentation(tmp_path):
    output = tmp_path / "existing.pptx"
    output.write_bytes(b"user-owned-content")
    with pytest.raises(ValueError, match="never overwritten"):
        concept.build(output)
    assert output.read_bytes() == b"user-owned-content"
